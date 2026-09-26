"""Отказы СК в Excel-своде (analytics_redesign_2026_09, задача 2.6)."""
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest

from .models import InsuranceCompany, InsuranceOffer, InsuranceSummary
from .services import company_statuses, get_excel_export_service


class DeclinedCompaniesExcelTests(TestCase):
    def setUp(self):
        request = InsuranceRequest.objects.create(
            client_name='ООО "СУ 888"', inn='2312100710', dfa_number='ТС-20496-ЛТ-КР',
            vehicle_info='Экскаватор LOVOL FR225E2-N', insurance_type='страхование спецтехники',
        )
        self.summary = InsuranceSummary.objects.create(request=request, status='ready')
        company_statuses.init_for_summary(self.summary)

    def offer(self, company, year=1, premium_2=None):
        return InsuranceOffer.objects.create(
            summary=self.summary, company_name=company, insurance_year=year,
            insurance_sum=Decimal('10000000'), franchise_1=Decimal('0'), premium_with_franchise_1=Decimal('250000'),
            franchise_2=Decimal('50000') if premium_2 else Decimal('0'), premium_with_franchise_2=premium_2,
        )

    def decline(self, *names):
        for name in names:
            company_statuses.set_status(self.summary, InsuranceCompany.objects.get(name=name).pk, 'declined')

    def sheet(self, client=False):
        data = get_excel_export_service().generate_summary_excel(self.summary, is_client_version=client)
        return load_workbook(BytesIO(data.getvalue())).worksheets[0]

    @staticmethod
    def rows(ws, first=10, last=30):
        return {r: {c.column_letter: c.value for c in ws[r] if c.value is not None} for r in range(first, last)}

    @staticmethod
    def merged(ws):
        return {str(m) for m in ws.merged_cells.ranges if m.min_row >= 10}

    def test_declined_rows_after_offers_simplified_template(self):
        self.offer('ВСК')
        self.decline('Согласие', 'Зетта')

        for client in (False, True):
            ws = self.sheet(client)
            rows = self.rows(ws)
            self.assertEqual(rows[10]['A'], 'ВСК')
            self.assertEqual(rows[11], {})  # разделитель
            self.assertEqual(rows[12], {'A': 'Зетта', 'B': 'Отказ от страхования'})  # по алфавиту
            self.assertEqual(rows[13], {'A': 'Согласие', 'B': 'Отказ от страхования'})
            self.assertEqual(rows[14], {})
            self.assertTrue({'B12:I12', 'B13:I13'} <= self.merged(ws))
            self.assertTrue(ws['B12'].font.italic)
            self.assertFalse(ws['A12'].font.bold)

    def test_full_template_with_multi_year_offers(self):
        self.offer('Абсолют', 1, premium_2=Decimal('240000'))
        self.offer('Абсолют', 2, premium_2=Decimal('230000'))
        self.offer('Альфа', 1)
        self.decline('Зетта')

        ws = self.sheet()
        rows = self.rows(ws)
        self.assertEqual((rows[10]['A'], rows[13]['A']), ('Абсолют', 'Альфа'))  # 11 — 2-й год, 12 — разделитель
        self.assertEqual(rows[14], {})
        self.assertEqual(rows[15], {'A': 'Зетта', 'B': 'Отказ от страхования'})
        self.assertIn('B15:O15', self.merged(ws))
        # в строке отказа нет формул и чисел — итоги и тарифы её не касаются
        self.assertFalse(any(isinstance(ws.cell(15, col).value, (int, float, Decimal)) for col in range(1, 18)))
        self.assertFalse(any(str(ws.cell(15, col).value or '').startswith('=') for col in range(1, 18)))

    def test_only_declined_companies(self):
        self.decline('Зетта', 'Абсолют')

        ws = self.sheet()
        rows = self.rows(ws)
        self.assertEqual(rows[10], {'A': 'Абсолют', 'B': 'Отказ от страхования'})  # формулы строки-образца убраны
        self.assertEqual(rows[11], {'A': 'Зетта', 'B': 'Отказ от страхования'})
        self.assertIn('B10:I10', self.merged(ws))

    def test_no_declined_companies_file_is_unchanged(self):
        self.offer('ВСК')
        company_statuses.set_undefined_to(self.summary, 'not_requested')

        ws = self.sheet()
        rows = self.rows(ws)
        self.assertEqual(rows[10]['A'], 'ВСК')
        self.assertTrue(all(not rows[r] for r in range(11, 30)))
        self.assertFalse(any(ws.cell(r, 2).value == 'Отказ от страхования' for r in range(10, 30)))

    def test_legacy_summary_without_statuses_keeps_note_only(self):
        legacy = InsuranceSummary.objects.create(
            request=InsuranceRequest.objects.create(
                client_name='ООО Старый', inn='7707083893', dfa_number='ТС-19000', vehicle_info='Тягач',
            ),
            status='ready', notes='СК Зетта - ОТКАЗ',
        )
        InsuranceOffer.objects.create(
            summary=legacy, company_name='ВСК', insurance_year=1, insurance_sum=Decimal('1000000'),
            franchise_1=Decimal('0'), premium_with_franchise_1=Decimal('50000'),
        )
        data = get_excel_export_service().generate_summary_excel(legacy)
        ws = load_workbook(BytesIO(data.getvalue())).worksheets[0]
        self.assertEqual(ws['C5'].value, 'СК Зетта - ОТКАЗ')
        self.assertTrue(all(ws.cell(r, 2).value != 'Отказ от страхования' for r in range(10, 30)))
