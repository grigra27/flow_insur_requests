"""Страница «Сотрудники» (analytics_redesign_2026_09, этап 4)."""
import uuid
from datetime import timedelta
from decimal import Decimal

from django.contrib.auth.models import Group, User
from django.contrib.contenttypes.models import ContentType
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from insurance_requests.models import InsuranceRequest
from summaries.models import InsuranceOffer, InsuranceSummary, StatusEvent
from summaries.services import analytics_employees as service


class EmployeesTestBase(TestCase):
    def setUp(self):
        admin_group, _ = Group.objects.get_or_create(name='Администраторы')
        self.users_group, _ = Group.objects.get_or_create(name='Пользователи')
        self.admin = User.objects.create_user(username='emp_admin', password='testpass123')
        self.admin.groups.add(admin_group)
        self.anna = self._employee('anna', 'Анна', 'Иванова')
        self.boris = self._employee('boris', 'Борис', 'Петров')
        self.reader = self._employee('reader', 'Вера', 'Смирнова')
        self.client.login(username='emp_admin', password='testpass123')
        self.now = timezone.now()
        self._counter = 0

    def _employee(self, username, first, last):
        user = User.objects.create_user(username=username, password='p', first_name=first, last_name=last)
        user.groups.add(self.users_group)
        return user

    def _request(self, user, days_ago, batch=None):
        self._counter += 1
        insurance_request = InsuranceRequest.objects.create(
            created_by=user, client_name=f'Клиент {self._counter}', inn='1234567890',
            insurance_type='КАСКО', dfa_number=f'DFA-{self._counter}', source_batch_id=batch,
        )
        InsuranceRequest.objects.filter(pk=insurance_request.pk).update(created_at=self.now - timedelta(days=days_ago))
        return insurance_request

    def _summary_with_offers(self, insurance_request, offers=2, days_ago=1):
        summary = InsuranceSummary.objects.create(request=insurance_request, status='collecting')
        for index in range(offers):
            offer = InsuranceOffer.objects.create(
                summary=summary, company_name=['Абсолют', 'Альфа', 'ВСК'][index], insurance_year=1,
                insurance_sum=Decimal('1000000'), premium_with_franchise_1=Decimal('10000'), franchise_1=Decimal('0'),
            )
            InsuranceOffer.objects.filter(pk=offer.pk).update(received_at=self.now - timedelta(days=days_ago))
        return summary

    def _status(self, summary, to_status, user, days_ago=1):
        event = StatusEvent.objects.create(
            content_type=ContentType.objects.get_for_model(InsuranceSummary), object_id=summary.pk,
            from_status='collecting', to_status=to_status, changed_by=user,
        )
        StatusEvent.objects.filter(pk=event.pk).update(changed_at=self.now - timedelta(days=days_ago))


class LoadBlockTests(EmployeesTestBase):
    def setUp(self):
        super().setUp()
        batch = uuid.uuid4()
        first = self._request(self.anna, 3, batch)
        self._request(self.anna, 3, batch)            # тот же файл
        self._request(self.anna, 10)
        self._request(self.anna, 120)                 # предыдущий 90-дневный период
        boris_request = self._request(self.boris, 5)
        self._request(self.boris, 200)                # вне обоих периодов

        summary = self._summary_with_offers(first, offers=3)
        self._summary_with_offers(boris_request, offers=2, days_ago=100)  # предложения вне периода
        self._status(summary, 'ready', self.anna)
        self._status(summary, 'sent', self.anna)
        self._status(summary, 'completed_accepted', self.boris)
        self._status(summary, 'completed_rejected', None)  # автозакрытие — без автора

    def _load(self, **params):
        return self.client.get(reverse('summaries:analytics_managers'), params).context['load']

    def test_rows_per_employee(self):
        rows = {row['name']: row for row in self._load()['rows']}

        anna = rows['Иванова Анна']
        self.assertEqual((anna['requests'], anna['files']), (3, 2))
        self.assertEqual(anna['offers'], 3)
        self.assertEqual((anna['assembled'], anna['sent'], anna['accepted']), (1, 1, 0))
        self.assertEqual(anna['share_pct'], Decimal('75'))
        self.assertEqual(anna['previous_requests'], 1)
        self.assertEqual(anna['change_pct'], Decimal('200'))

        boris = rows['Петров Борис']
        self.assertEqual((boris['requests'], boris['offers'], boris['accepted']), (1, 0, 1))
        self.assertIsNone(boris['change_pct'])

    def test_readers_are_listed_last(self):
        rows = self._load()['rows']

        self.assertEqual(rows[-1]['name'], 'Смирнова Вера')
        self.assertTrue(rows[-1]['is_reader'])
        self.assertNotIn('emp_admin', [row['name'] for row in rows])

    def test_totals_exclude_auto_close(self):
        totals = self._load()['totals']

        self.assertEqual((totals['requests'], totals['files'], totals['offers']), (4, 3, 3))
        self.assertEqual((totals['accepted'], totals['rejected']), (1, 0))
        self.assertEqual(totals['per_week'], Decimal(4) / (Decimal(90) / Decimal(7)))

    def test_longer_period_and_monthly_chart(self):
        load = self._load(period='365')

        self.assertEqual(load['totals']['requests'], 6)
        self.assertEqual(load['chart']['granularity'], 'month')
        self.assertEqual(sum(sum(series['data']) for series in load['chart']['series']), 6)

    def test_weekly_chart_and_stable_colors(self):
        load = self._load()

        self.assertEqual(load['chart']['granularity'], 'week')
        colors = {series['label']: series['color'] for series in load['chart']['series']}
        # Цвет закреплён за сотрудником (по порядку id), не за местом в таблице.
        self.assertEqual(colors['Иванова Анна'], service.SERIES_COLORS[0])
        self.assertEqual(colors['Петров Борис'], service.SERIES_COLORS[1])

    def test_custom_dates(self):
        start = (timezone.localdate() - timedelta(days=6)).isoformat()
        load = self._load(start_date=start)

        self.assertEqual(load['totals']['requests'], 3)

    def test_page_renders_for_admin_only(self):
        response = self.client.get(reverse('summaries:analytics_managers'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'summaries/analytics_employees.html')
        self.assertContains(response, 'Нагрузка за период')
        self.assertContains(response, 'только просмотр')

        self.client.login(username='anna', password='p')
        self.assertEqual(self.client.get(reverse('summaries:analytics_managers')).status_code, 403)


class PresenceBlockTests(EmployeesTestBase):
    def _day(self, user, days_ago, start, end, minutes, **fields):
        from datetime import datetime, time
        from summaries.models import UserDailyActivity

        day = timezone.localdate() - timedelta(days=days_ago)
        tz = timezone.get_default_timezone()
        return UserDailyActivity.objects.create(
            user=user, date=day,
            first_seen_at=timezone.make_aware(datetime.combine(day, time(*start)), tz),
            last_seen_at=timezone.make_aware(datetime.combine(day, time(*end)), tz),
            active_minutes=minutes, **fields,
        )

    def test_presence_rows(self):
        # Найти ближайшую субботу в пределах периода для проверки выходных.
        saturday_ago = (timezone.localdate().weekday() - 5) % 7 or 7
        self._day(self.anna, 1, (9, 0), (17, 0), 200, logins_count=1, page_views=40, has_request_data=True,
                  hourly={'9': 10, '10': 5})
        self._day(self.anna, 2, (10, 0), (21, 30), 100, logins_count=1, crud_actions=5)
        self._day(self.anna, saturday_ago + 7, (11, 0), (12, 0), 30, page_views=3, has_request_data=True)
        self._day(self.reader, 3, (12, 0), (12, 0), 0, logins_count=1)  # вошла, ничего не делала

        presence = self.client.get(reverse('summaries:analytics_managers')).context['presence']
        rows = {row['name']: row for row in presence['rows']}

        anna = rows['Иванова Анна']
        self.assertEqual((anna['login_days'], anna['active_days']), (2, 3))
        self.assertEqual(anna['typical_start'], '10:00')
        self.assertEqual(anna['typical_end'], '17:00')
        self.assertEqual(anna['active_minutes_per_day'], 110)
        self.assertEqual(anna['active_hours_total'], Decimal(330) / Decimal(60))
        self.assertEqual(anna['late_days'], 1)
        self.assertEqual(anna['weekend_days'], 1)
        self.assertEqual(anna['days_without_views'], 1)

        reader = rows['Смирнова Вера']
        self.assertEqual((reader['login_days'], reader['active_days']), (1, 0))
        self.assertIsNone(reader['typical_start'])

        day_index = (timezone.localdate() - timedelta(days=1)).weekday()
        self.assertEqual(presence['heatmap']['team'][day_index][9], 10)
        self.assertEqual(presence['heatmap']['employees'][str(self.anna.pk)][day_index][10], 5)

    def test_coverage_note_when_data_starts_late(self):
        self._day(self.anna, 5, (9, 0), (10, 0), 60, page_views=1, has_request_data=True)

        response = self.client.get(reverse('summaries:analytics_managers'))

        self.assertTrue(response.context['presence']['coverage_partial'])
        self.assertContains(response, 'Присутствие')
        self.assertContains(response, 'Когда работают')


class SpeedBlockTests(EmployeesTestBase):
    def _event(self, model, object_id, to_status, moment, user=None):
        event = StatusEvent.objects.create(
            content_type=ContentType.objects.get_for_model(model), object_id=object_id,
            from_status='', to_status=to_status, changed_by=user,
        )
        StatusEvent.objects.filter(pk=event.pk).update(changed_at=moment)

    def test_medians_and_attribution(self):
        request = self._request(self.anna, 5)
        created = InsuranceRequest.objects.get(pk=request.pk).created_at
        self._event(InsuranceRequest, request.pk, 'emails_sent', created + timedelta(hours=2), self.anna)

        summary = self._summary_with_offers(request, offers=2, days_ago=4)
        last_offer = self.now - timedelta(days=4)
        self._event(InsuranceSummary, summary.pk, 'sent', last_offer + timedelta(hours=6), self.boris)
        InsuranceOffer.objects.create(  # внесено после отправки — не считается «последним до отправки»
            summary=summary, company_name='ВСК', insurance_year=1, insurance_sum=Decimal('1'),
            premium_with_franchise_1=Decimal('1'), franchise_1=Decimal('0'),
        )
        InsuranceSummary.objects.filter(pk=summary.pk).update(
            status='completed_accepted', completed_at=last_offer + timedelta(hours=6 + 48),
        )

        speed = self.client.get(reverse('summaries:analytics_managers')).context['speed']
        anna = next(row for row in speed['rows'] if row['user_id'] == self.anna.pk)

        self.assertAlmostEqual(anna['upload_to_emails']['median_hours'], 2.0, places=3)
        self.assertEqual(anna['upload_to_emails']['median'], '2,0 ч')
        # Атрибуция — владелец заявки (Анна), хотя отправлял Борис.
        self.assertAlmostEqual(anna['offer_to_client']['median_hours'], 6.0, places=3)
        self.assertEqual(anna['client_wait']['median'], '2,0 дн')
        self.assertEqual(speed['team']['offer_to_client']['count'], 1)

    def test_stuck_summaries(self):
        old = self._request(self.boris, 20)
        summary = InsuranceSummary.objects.create(request=old, status='collecting')
        InsuranceSummary.objects.filter(pk=summary.pk).update(created_at=self.now - timedelta(days=10))
        fresh = self._request(self.boris, 2)
        InsuranceSummary.objects.create(request=fresh, status='ready')

        response = self.client.get(reverse('summaries:analytics_managers'))
        speed = response.context['speed']

        self.assertEqual([item['summary_id'] for item in speed['stuck']], [summary.pk])
        self.assertEqual(speed['stuck'][0]['age_days'], 10)
        boris = next(row for row in speed['rows'] if row['user_id'] == self.boris.pk)
        self.assertEqual(boris['stuck'], 1)
        self.assertContains(response, 'Скорость на своих этапах')
        self.assertContains(response, 'Зависшие своды')

    def test_format_hours(self):
        self.assertEqual(service.format_hours(None), None)
        self.assertEqual(service.format_hours(0.25), '15 мин')
        self.assertEqual(service.format_hours(5.24), '5,2 ч')
        self.assertEqual(service.format_hours(84), '3,5 дн')


class DossierAndExportTests(EmployeesTestBase):
    def setUp(self):
        super().setUp()
        from datetime import datetime, time
        from summaries.models import UserDailyActivity

        self.request = self._request(self.anna, 3)
        self._summary_with_offers(self.request, offers=2)
        day = timezone.localdate() - timedelta(days=1)
        tz = timezone.get_default_timezone()
        UserDailyActivity.objects.create(
            user=self.anna, date=day, logins_count=1, page_views=12, form_actions=3, crud_actions=4,
            sections={'summaries': 10, 'requests': 5}, hourly={'10': 7}, active_minutes=95,
            first_seen_at=timezone.make_aware(datetime.combine(day, time(9, 30)), tz),
            last_seen_at=timezone.make_aware(datetime.combine(day, time(16, 5)), tz),
            has_request_data=True,
        )

    def test_dossier_blocks(self):
        response = self.client.get(reverse('summaries:analytics_manager_detail', args=[self.anna.pk]))

        self.assertEqual(response.status_code, 200)
        context = response.context
        self.assertEqual(context['name'], 'Иванова Анна')
        self.assertEqual((context['load_row']['requests'], context['load_row']['offers']), (1, 2))
        self.assertEqual(context['presence_row']['active_days'], 1)
        self.assertEqual(context['day_rows'][0]['sections'][0], {'label': 'Своды', 'count': 10})
        day_index = (timezone.localdate() - timedelta(days=1)).weekday()
        self.assertEqual(context['heatmap']['grid'][day_index][10], 7)
        for text in ('Нагрузка за период', 'Присутствие', 'Скорость на своих этапах', 'Лента действий', 'js-tl-filter'):
            self.assertContains(response, text)
        StatusEvent.objects.create(
            content_type=ContentType.objects.get_for_model(InsuranceSummary),
            object_id=InsuranceSummary.objects.get(request=self.request).pk,
            from_status='collecting', to_status='ready', changed_by=self.anna,
        )
        response = self.client.get(reverse('summaries:analytics_manager_detail', args=[self.anna.pk]))
        self.assertContains(response, '«Сбор предложений» → «Готов к отправке»')
        for removed in ('Радар', 'Quality', 'Любимые СК'):
            self.assertNotContains(response, removed)

    def test_dossier_for_reader_and_unknown_user(self):
        reader = self.client.get(reverse('summaries:analytics_manager_detail', args=[self.reader.pk]))
        self.assertEqual(reader.status_code, 200)
        self.assertTrue(reader.context['load_row']['is_reader'])

        missing = self.client.get(reverse('summaries:analytics_manager_detail', args=[99999]))
        self.assertEqual(missing.status_code, 404)
        self.assertContains(missing, 'Сотрудник не найден', status_code=404)

    def test_overview_export_sheets(self):
        from io import BytesIO
        from openpyxl import load_workbook

        response = self.client.get(reverse('summaries:export_analytics_managers_widget'), {'period': '180'})

        self.assertEqual(response.status_code, 200)
        self.assertIn('employees_', response['Content-Disposition'])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Нагрузка', 'Присутствие', 'Скорость', 'Зависшие своды'])
        self.assertEqual(workbook['Нагрузка']['A6'].value, 'Иванова Анна')

    def test_dossier_export_sheets(self):
        from io import BytesIO
        from openpyxl import load_workbook

        response = self.client.get(reverse('summaries:export_analytics_manager_detail', args=[self.anna.pk]))

        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Сводка', 'По дням'])
        self.assertEqual(workbook['По дням']['D6'].value, 95)
