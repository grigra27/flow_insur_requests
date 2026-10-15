"""Блок «Риски ОСВР» — только филиал «Краснодар» (решение владельца 2026-10-15).

Обязателен для краснодарских заявок с даты выката (OSVR_SINCE); в своде V2 — колонка «ОСВР», ограничения —
в комментариях; второй вариант «с ОСВР» подписан над ценником. История — командой osvr_from_notes.
"""
import json
from decimal import Decimal
from io import BytesIO, StringIO

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.templates import EmailTemplateGenerator
from insurance_requests.models import InsuranceRequest
from summaries.management.commands.osvr_from_notes import parse
from summaries.models import InsuranceOffer, InsuranceSummary, InsurerResponse
from summaries.response_sections import OSVR, OSVR_SINCE, required_sections, shown_sections
from summaries.response_template import build_v2
from summaries.services.insurer_response import display_lines
from summaries.services.summary_export_v2 import SummaryExportV2Service

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
AFTER = timezone.make_aware(timezone.datetime(OSVR_SINCE.year, OSVR_SINCE.month, OSVR_SINCE.day, 12, 0))
BEFORE = timezone.make_aware(timezone.datetime(2026, 10, 10, 12, 0))


class _Request:
    def __init__(self, branch, created_at=None):
        self.branch, self.created_at, self.insurance_type = branch, created_at, 'КАСКО'


class _Response:
    def __init__(self, osvr_status=''):
        self.osvr_status = osvr_status


class OsvrRuleTests(SimpleTestCase):
    def test_required_only_for_krasnodar_since_rollout(self):
        self.assertTrue(OSVR.is_required(_Request('Краснодар', AFTER)))
        self.assertTrue(OSVR.is_required(_Request('Краснодар', None)))
        self.assertFalse(OSVR.is_required(_Request('Краснодар', BEFORE)))
        self.assertFalse(OSVR.is_required(_Request('Москва', AFTER)))
        self.assertEqual([s.key for s in required_sections(_Request('Краснодар', AFTER))], ['inspection', 'osvr'])

    def test_shown_when_answer_exists(self):
        old = _Request('Краснодар', BEFORE)
        self.assertNotIn('osvr', [s.key for s in shown_sections(old, [_Response()])])
        self.assertIn('osvr', [s.key for s in shown_sections(old, [_Response(), _Response('none')])])

    def test_labels_without_commas(self):
        for label in OSVR.fields[0].choice_labels:
            self.assertNotIn(',', label)

    def test_parse_notes_seen_on_prod(self):
        cases = {
            'без ОСВР': ('none', ''), 'Без ОСВР. Согласование РНПК.': ('none', ''),
            'Риски ОСВР не включены. В расчет включены риски по событиям': ('none', ''),
            'с ОСВР': ('included', ''), 'с рисками ОСВР': ('included', ''),
            'С ОСВР (Без риска БПЛА)': ('limited', 'Без риска БПЛА'),
            'с ОСВР ( за исключением работы ПВО)': ('limited', 'за исключением работы ПВО'),
            '2 вариант с ОСВР': ('variant_2', ''),
            'Бой стекол входит в покрытие. 2 вариант - с действиями животных. без ОСВР': ('none', ''),
            'вариант 2 с ОСВР. Ограничения по ОСВР: - ДНР; - ЛНР. вариант 2 с ОСВР. Ограничения по ОСВР: - ДНР; - ЛНР.':
                ('variant_2', '- ДНР; - ЛНР.'),
            'с ОСВР, без ОСВР': (None, ''), 'Требуется осмотр': (None, ''),
        }
        for text, expected in cases.items():
            self.assertEqual(parse(text), expected, text)


def _filled(workbook, blocks, company='Согаз', premium_2=None):
    ws = workbook['Ответ СК']
    ws['B2'], ws['B3'] = company, 'Российская Федерация'
    for col, value in zip('BDEF', (7_300_000, 200_000, 0, 1)):
        ws[f'{col}6'] = value
    if premium_2:
        ws['H6'], ws['I6'], ws['J6'] = premium_2, 0, 1
    for name, value in blocks.items():
        sheet, ref = next(iter(workbook.defined_names[name].destinations))
        workbook[sheet][ref.replace('$', '')] = value
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile('ответ.xlsx', buffer.getvalue(), content_type=XLSX)


@override_settings(RESPONSE_TEMPLATE_V2_FOR_ALL=True, RESPONSE_TEMPLATE_V2_SINCE='2026-10-08')
class OsvrFlowTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.staff = User.objects.create_user('staff', password='p')
        cls.staff.groups.add(Group.objects.get_or_create(name='Пользователи')[0])
        cls.request_obj = InsuranceRequest.objects.create(
            client_name='ООО Юг', inn='1', dfa_number='ТС-20928', object_description='VOYAH DREAM',
            insurance_type='КАСКО', branch='Краснодар', condition='new')
        InsuranceRequest.objects.filter(pk=cls.request_obj.pk).update(created_at=AFTER)
        cls.request_obj.refresh_from_db()
        cls.summary = InsuranceSummary.objects.create(request=cls.request_obj, status='collecting')

    def _upload(self, upload):
        self.client.force_login(self.staff)
        response = self.client.post(reverse('summaries:upload_multiple_company_responses', args=[self.summary.pk]),
                                    {'excel_files': [upload]})
        return json.loads(response.content)['results'][0]

    def _template(self):
        return load_workbook(BytesIO(build_v2(self.request_obj, self.summary.pk)))

    def _sheet(self):
        data = SummaryExportV2Service().generate(self.summary, is_client_version=True).getvalue()
        return load_workbook(BytesIO(data))['Свод']

    def test_template_has_block_with_dropdown(self):
        wb = self._template()
        self.assertIn('resp_osvr_status', wb.defined_names)
        self.assertIn('resp_osvr_terms', wb.defined_names)
        texts = ' '.join(str(c.value) for row in wb['Ответ СК'].iter_rows() for c in row if c.value)
        self.assertIn('Риски ОСВР', texts)
        lists = [dv.formula1 for dv in wb['Ответ СК'].data_validations.dataValidation if dv.type == 'list']
        self.assertIn('"Включены,Включены с ограничениями,Только во втором варианте,Не включены"', lists)

    def test_answers_in_summary_v2_column_notes_and_variant_2_caption(self):
        ok = self._upload(_filled(self._template(), {
            'resp_inspection_status': 'Осмотр не требуется', 'resp_osvr_status': 'Не включены'}))
        self.assertTrue(ok['success'], ok)
        ok = self._upload(_filled(self._template(), {
            'resp_inspection_status': 'Осмотр не требуется', 'resp_osvr_status': 'Только во втором варианте',
            'resp_osvr_terms': 'кроме ДНР, ЛНР'}, company='Альфа', premium_2=245_000))
        self.assertTrue(ok['success'], ok)
        alfa = InsurerResponse.objects.get(summary=self.summary, company_name='Альфа')
        self.assertEqual((alfa.osvr_status, alfa.osvr_terms), ('variant_2', 'кроме ДНР, ЛНР'))

        ws = self._sheet()
        titles = [ws.cell(row=9, column=c).value for c in range(1, ws.max_column + 1)]
        col = titles.index('ОСВР') + 1
        rows = {ws.cell(row=r, column=1).value: r for r in range(11, ws.max_row + 1) if ws.cell(row=r, column=1).value}
        self.assertEqual(ws.cell(row=rows['Альфа'], column=col).value, 'Во 2-м варианте')
        self.assertEqual(ws.cell(row=rows['Согаз'], column=col).value, 'Не включены')
        notes = ws.cell(row=rows['Альфа'], column=titles.index('Комментарии') + 1).value
        self.assertIn('ОСВР: кроме ДНР, ЛНР', notes)
        self.assertEqual(ws['C10'].value, 'Предложение 2 · с ОСВР')  # второй вариант — только у Альфы, с ОСВР

        line = next(l for l in display_lines(self.summary, ['Альфа'])['Альфа']['lines'] if l['label'] == 'ОСВР')
        self.assertEqual(line['text'], 'Только во втором варианте · кроме ДНР, ЛНР')

    def test_missing_answer_rejected(self):
        result = self._upload(_filled(self._template(), {'resp_inspection_status': 'Осмотр не требуется'}))
        self.assertFalse(result['success'])
        self.assertIn('ОСВР', result['error_message'])

    def test_email_line_only_for_krasnodar(self):
        generator = EmailTemplateGenerator()
        body = generator.generate_email_body(self.request_obj.to_dict())
        self.assertIn('включены ли в расчёт риски ОСВР', body)
        self.assertNotIn('ОСВР', generator.generate_email_body(self.request_obj.to_dict() | {'branch': 'Москва'}))


class OsvrHistoryCommandTests(TestCase):
    def setUp(self):
        request = InsuranceRequest.objects.create(client_name='ООО Юг', inn='1', dfa_number='ТС-20928',
                                                  object_description='VOYAH DREAM', branch='Краснодар')
        InsuranceRequest.objects.filter(pk=request.pk).update(created_at=BEFORE)  # до блока ОСВР
        self.summary = InsuranceSummary.objects.create(request=request, status='ready')
        for company, notes in (('Альфа', 'с ОСВР'), ('Ингосстрах', 'Без ОСВР. Россия, СНГ'), ('Пари', 'с ОСВР, без ОСВР')):
            InsuranceOffer.objects.create(summary=self.summary, company_name=company, insurance_year=1,
                                          insurance_sum=Decimal('1000000'), franchise_1=Decimal('0'),
                                          premium_with_franchise_1=Decimal('50000'))
            InsurerResponse.objects.filter(summary=self.summary, company_name=company).update(notes=notes)

    def _run(self, *args):
        out = StringIO()
        call_command('osvr_from_notes', *args, stdout=out)
        return out.getvalue()

    def test_report_then_apply_and_old_summary_shows_column(self):
        report = self._run()
        self.assertIn("к заполнению: 2 {'included': 1, 'none': 1}", report)
        self.assertIn('Не разобрано (оставлено пустым): 1', report)
        self.assertFalse(InsurerResponse.objects.exclude(osvr_status='').exists())
        self._run('--apply')
        statuses = dict(InsurerResponse.objects.filter(summary=self.summary).values_list('company_name', 'osvr_status'))
        self.assertEqual(statuses, {'Альфа': 'included', 'Ингосстрах': 'none', 'Пари': ''})
        # старый свод (блок не был нужен) — колонка появляется, раз ответ есть
        ws = load_workbook(BytesIO(SummaryExportV2Service().generate(self.summary, True).getvalue()))['Свод']
        self.assertIn('ОСВР', [ws.cell(row=9, column=c).value for c in range(1, ws.max_column + 1)])
        self.assertIn("к заполнению: 0", self._run())
