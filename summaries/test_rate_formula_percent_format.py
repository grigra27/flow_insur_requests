"""
Регрессия: тариф в сводах показывался как 200,00% вместо 2,00%.

Ячейки тарифа (E, K) в шаблонах имеют процентный формат `0.00%`, который сам
умножает значение на 100 при отображении. Поэтому формула должна возвращать
долю (премия / страховая сумма), а не проценты: `*100` в формуле при
процентном формате даёт завышение в 100 раз.
"""

from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from insurance_requests.models import InsuranceRequest
from summaries.models import InsuranceCompany, InsuranceOffer, InsuranceSummary
from summaries.services.excel_services import ExcelExportService


TEMPLATES_DIR = Path(settings.BASE_DIR) / 'templates'
SUMMARY_TEMPLATES = (
    'summary_template.xlsx',
    'client_summary_template.xlsx',
    'summary_template_simplified.xlsx',
    'client_summary_template_simplified.xlsx',
)
RATE_COLUMNS = ('E', 'K')


class RateFormulaAssertionsMixin:
    def assertRateCellConsistent(self, cell, where):
        formula = cell.value
        if not (isinstance(formula, str) and formula.startswith('=')):
            return
        normalized = formula.replace(' ', '')
        if '%' in (cell.number_format or ''):
            self.assertNotIn(
                '*100',
                normalized,
                f'{where}: формула тарифа умножает на 100 при процентном '
                f'формате ячейки ({cell.number_format}) — тариф будет '
                f'завышен в 100 раз: {formula}',
            )


class RateTemplatesInvariantTests(RateFormulaAssertionsMixin, TestCase):
    def test_template_rate_cells_do_not_double_percent(self):
        for filename in SUMMARY_TEMPLATES:
            sheet = load_workbook(TEMPLATES_DIR / filename)['summary_template_sheet']
            for column in RATE_COLUMNS:
                with self.subTest(template=filename, cell=f'{column}10'):
                    self.assertRateCellConsistent(sheet[f'{column}10'], f'{filename}!{column}10')

    def test_template_rate_formula_uses_percent_format(self):
        for filename in SUMMARY_TEMPLATES:
            cell = load_workbook(TEMPLATES_DIR / filename)['summary_template_sheet']['E10']
            with self.subTest(template=filename):
                self.assertEqual(cell.value, '=IF(AND(C10<>0,F10<>0),F10/C10,"")')
                self.assertIn('%', cell.number_format)

    def test_fallback_rate_formula_is_consistent_with_percent_format(self):
        template_path = TEMPLATES_DIR / 'summary_template.xlsx'
        service = ExcelExportService(str(template_path))
        sheet = Workbook().active
        for column in RATE_COLUMNS:
            with self.subTest(column=column):
                sheet[f'{column}11'].number_format = '0.00%'
                service._create_rate_formula(sheet, 11, column)
                self.assertRateCellConsistent(sheet[f'{column}11'], f'fallback {column}11')


class RateExportInvariantTests(RateFormulaAssertionsMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        InsuranceCompany.objects.get_or_create(
            name='Абсолют',
            defaults={'display_name': 'Абсолют', 'sort_order': 1, 'is_active': True},
        )
        request = InsuranceRequest.objects.create(
            client_name='ООО Тариф',
            inn='1234567890',
            dfa_number='ТС-20276-ЛА-ПС',
            vehicle_info='Haval F7X',
        )
        cls.summary = InsuranceSummary.objects.create(request=request, status='ready')
        for year, insurance_sum in ((1, '3549000'), (2, '3016650'), (3, '2564153')):
            InsuranceOffer.objects.create(
                summary=cls.summary,
                company_name='Абсолют',
                insurance_year=year,
                insurance_sum=Decimal(insurance_sum),
                franchise_1=Decimal('0'),
                premium_with_franchise_1=Decimal('70980'),
                franchise_2=Decimal('50000'),
                premium_with_franchise_2=Decimal('60000'),
            )

    def _assert_all_rate_cells(self, label):
        service = ExcelExportService(str(TEMPLATES_DIR / 'summary_template.xlsx'))
        for is_client_version in (False, True):
            workbook = load_workbook(
                service.generate_summary_excel(self.summary, is_client_version=is_client_version)
            )
            sheet = workbook['summary_template_sheet']
            checked = 0
            for row in range(10, sheet.max_row + 1):
                for column in RATE_COLUMNS:
                    cell = sheet[f'{column}{row}']
                    if isinstance(cell.value, str) and cell.value.startswith('='):
                        checked += 1
                        self.assertRateCellConsistent(
                            cell, f'{label} client={is_client_version} {column}{row}'
                        )
            self.assertGreaterEqual(checked, 3, f'{label} client={is_client_version}')

    def test_full_exports_rate_cells(self):
        self._assert_all_rate_cells('full')

    def test_simplified_exports_rate_cells(self):
        InsuranceOffer.objects.filter(summary=self.summary).update(
            premium_with_franchise_2=None,
            franchise_2=None,
        )
        self._assert_all_rate_cells('simplified')
