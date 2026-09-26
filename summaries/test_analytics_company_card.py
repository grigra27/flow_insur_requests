"""Карточка страховой компании (analytics_redesign_2026_09, этап 3)."""
from decimal import Decimal
from io import BytesIO
from urllib.parse import quote

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest

from .models import InsuranceCompany, InsuranceOffer, InsuranceSummary
from .services import company_statuses
from .services.analytics_insurance_company_card import build_card_payload, dfa_kind
from .views import _build_deal_price_row

FILTERS = {'start_date': None, 'end_date': None, 'branch': '', 'insurance_type': ''}


def make_deal(dfa, winner, prices, branch='Краснодар', insurance_type='КАСКО', client='ООО Клиент', status='completed_accepted'):
    request = InsuranceRequest.objects.create(
        client_name=client, inn='7707083893', dfa_number=dfa, branch=branch, insurance_type=insurance_type,
        vehicle_info='Тягач', manufacturing_year='2024', condition='new',
    )
    summary = InsuranceSummary.objects.create(request=request, status=status, selected_company=winner,
                                              selected_franchise_variant=1 if winner else None)
    for company, premium in prices.items():
        InsuranceOffer.objects.create(
            summary=summary, company_name=company, insurance_year=1, insurance_sum=Decimal('10000000'),
            franchise_1=Decimal('0'), premium_with_franchise_1=Decimal(premium),
        )
    return summary


class CompanyCardPayloadTests(TestCase):
    def setUp(self):
        make_deal('ТС-1-ГА-КР', 'Абсолют', {'Абсолют': 100000, 'Зетта': 120000})               # Абсолют дешевле, выиграл
        make_deal('ТС-2-ЛА-МСК', 'Зетта', {'Абсолют': 90000, 'Зетта': 110000}, branch='Москва')  # Абсолют дешевле, проиграл
        make_deal('ТС-3-ЛТ-КР', 'Зетта', {'Зетта': 50000, 'Альфа': 60000},
                  insurance_type='страхование спецтехники')                                     # Абсолют не участвовал

    def payload(self, company='Абсолют'):
        return build_card_payload(company, FILTERS, _build_deal_price_row)

    def test_volume_and_shares(self):
        data = self.payload()
        self.assertEqual(data['kpi']['deals'], 1)
        self.assertEqual(data['kpi']['market_deals'], 3)
        self.assertEqual(round(data['kpi']['deals_share_pct']), 33)
        self.assertEqual(data['kpi']['insured_sum'], Decimal('10000000'))
        self.assertEqual(data['kpi']['tariff_pct'], Decimal('1'))  # 100 000 / 10 000 000

    def test_price_block(self):
        price = self.payload()['price']
        self.assertEqual((price['participated'], price['won'], price['lost']), (2, 1, 1))
        self.assertEqual((price['won_cheapest'], price['won_not_cheapest'], price['cheapest_lost']), (1, 0, 1))
        self.assertEqual(price['avg_rank'], Decimal('1'))
        self.assertEqual(self.payload()['cheapest_lost_to'], [('Зетта', 1)])

        zetta = self.payload('Зетта')['price']
        self.assertEqual((zetta['won_cheapest'], zetta['won_not_cheapest']), (1, 1))  # 2-я сделка — не ценой

    def test_dimensions_and_deal_list(self):
        data = self.payload('Зетта')
        dims = dict(data['dimensions'])
        self.assertEqual({row['label']: row['deals'] for row in dims['Вид лизинга (код ДФА)']},
                         {'Легковые (ЛА)': 1, 'Спецтехника (ЛТ)': 1})
        self.assertEqual({row['label'] for row in dims['Филиалы']}, {'Москва', 'Краснодар'})
        self.assertEqual(len(data['deal_rows']), 3)
        lost = [row for row in data['deal_rows'] if not row['won']][0]
        self.assertEqual((lost['winner'], lost['rank'], lost['compared']), ('Абсолют', 2, 2))

    def test_dfa_kind_fallbacks(self):
        self.assertEqual(dfa_kind(InsuranceRequest(dfa_number='ТС-20762-ЛА')), 'Легковые (ЛА)')
        self.assertEqual(dfa_kind(InsuranceRequest(dfa_number='2026', insurance_type='страхование спецтехники')),
                         'Спецтехника (ЛТ)')
        self.assertEqual(dfa_kind(InsuranceRequest(dfa_number='ТС-1', insurance_type='КАСКО')), 'Не указан')

    def test_funnel_from_company_statuses(self):
        self.assertFalse(self.payload()['funnel']['has_data'])

        summary = make_deal('ТС-9-ГА-КР', 'Абсолют', {'Абсолют': 80000}, status='collecting')
        company_statuses.init_for_summary(summary)
        company_statuses.set_status(summary, InsuranceCompany.objects.get(name='Зетта').pk, 'declined')
        other = make_deal('ТС-10-ГА-КР', '', {}, status='collecting')
        company_statuses.init_for_summary(other)
        company_statuses.set_status(other, InsuranceCompany.objects.get(name='Зетта').pk, 'requested')
        InsuranceSummary.objects.filter(pk=summary.pk).update(status='completed_accepted')

        absolut = self.payload()['funnel']
        self.assertEqual((absolut['asked'], absolut['offered'], absolut['won']), (1, 1, 1))
        zetta = self.payload('Зетта')['funnel']
        self.assertEqual((zetta['asked'], zetta['declined'], zetta['no_answer']), (2, 1, 1))
        self.assertEqual(zetta['declined_pct'], Decimal('50'))
        self.assertEqual(zetta['by_type'][0]['label'], 'КАСКО')

    def test_legacy_summary_counts_only_when_statuses_were_set_by_hand(self):
        legacy = make_deal('ТС-11-ГА-КР', '', {'Альфа': 70000}, status='sent')  # «предложение» появилось само
        self.assertFalse(self.payload('Альфа')['funnel']['has_data'])

        company_statuses.set_status(legacy, InsuranceCompany.objects.get(name='Зетта').pk, 'declined')
        self.assertEqual(self.payload('Альфа')['funnel']['offered'], 1)
        self.assertEqual(self.payload('Зетта')['funnel']['declined'], 1)


class CompanyCardViewTests(TestCase):
    def setUp(self):
        admin = User.objects.create_user(username='card_admin', password='pwd')
        admin.groups.add(Group.objects.get_or_create(name='Администраторы')[0])
        self.client.login(username='card_admin', password='pwd')
        make_deal('ТС-1-ГА-КР', 'Абсолют', {'Абсолют': 100000, 'Зетта': 120000})

    def url(self, company, name='summaries:analytics_insurance_company_card'):
        return reverse(name, args=[company])

    def test_card_renders_and_links_from_list(self):
        response = self.client.get(self.url('Абсолют'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Почему выигрывает и проигрывает')
        self.assertContains(response, 'Воронка')

        listing = self.client.get(reverse('summaries:analytics_insurance_companies'))
        self.assertContains(listing, self.url('Абсолют'))

    def test_unknown_company_is_404(self):
        self.assertEqual(self.client.get(self.url('Нет такой')).status_code, 404)

    def test_export(self):
        response = self.client.get(self.url('Абсолют', 'summaries:export_analytics_insurance_company_card'))
        self.assertEqual(response.status_code, 200)
        self.assertIn(quote('company_Абсолют'), response['Content-Disposition'])
        workbook = load_workbook(BytesIO(response.content))
        self.assertIn('Сделки', workbook.sheetnames)
        self.assertEqual(workbook['Сводка']['B6'].value, 1)

    def test_regular_user_has_no_access(self):
        user = User.objects.create_user(username='card_user', password='pwd')
        user.groups.add(Group.objects.get_or_create(name='Пользователи')[0])
        self.client.login(username='card_user', password='pwd')
        self.assertNotEqual(self.client.get(self.url('Абсолют')).status_code, 200)
