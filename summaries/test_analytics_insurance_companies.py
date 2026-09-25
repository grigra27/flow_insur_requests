from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest
from summaries.models import InsuranceOffer, InsuranceSummary


class InsuranceCompaniesAnalyticsTests(TestCase):
    def setUp(self):
        self.client = Client()

        self.admin_group, _ = Group.objects.get_or_create(name='Администраторы')
        self.user_group, _ = Group.objects.get_or_create(name='Пользователи')

        self.admin_user = User.objects.create_user(
            username='analytics_companies_admin',
            email='analytics_companies_admin@example.com',
            password='testpass123',
            first_name='Алексей',
            last_name='Петров',
        )
        self.admin_user.groups.add(self.admin_group)

        self.regular_user = User.objects.create_user(
            username='analytics_companies_user',
            email='analytics_companies_user@example.com',
            password='testpass123',
        )
        self.regular_user.groups.add(self.user_group)

        self.client.login(username='analytics_companies_admin', password='testpass123')

        now = timezone.now()

        self.summary_moscow = self._create_completed_summary(
            dfa_number='DFA-COMP-001',
            branch='Москва',
            insurance_type='КАСКО',
            manager_name='Иванов И.И.',
            deal_status='new',
            selected_company='Абсолют',
            selected_variant=1,
            response_deadline=now - timedelta(days=7),
            request_created_at=now - timedelta(days=11),
            summary_created_at=now - timedelta(days=10),
            completed_at=now - timedelta(days=8),
        )
        self._add_offer(self.summary_moscow, 'Абсолют', '10000.00', received_at=now - timedelta(days=9))
        self._add_offer(self.summary_moscow, 'Альфа', '9000.00', received_at=now - timedelta(days=9))
        self._add_offer(self.summary_moscow, 'ВСК', '12000.00', received_at=now - timedelta(days=9))

        self.summary_spb = self._create_completed_summary(
            dfa_number='DFA-COMP-002',
            branch='Санкт-Петербург',
            insurance_type='страхование имущества',
            manager_name='',
            deal_status='prolongation',
            selected_company='Альфа',
            selected_variant=None,
            response_deadline=now - timedelta(days=4),
            request_created_at=now - timedelta(days=7),
            summary_created_at=now - timedelta(days=6),
            completed_at=now - timedelta(days=3),
        )
        self._add_offer(self.summary_spb, 'Альфа', '8000.00', received_at=now - timedelta(days=5))
        self._add_offer(self.summary_spb, 'Абсолют', '8500.00', received_at=now - timedelta(days=5))

    def _create_completed_summary(
        self,
        *,
        dfa_number,
        branch,
        insurance_type,
        manager_name,
        deal_status,
        selected_company,
        selected_variant,
        response_deadline,
        request_created_at,
        summary_created_at,
        completed_at,
    ):
        insurance_request = InsuranceRequest.objects.create(
            created_by=self.admin_user,
            client_name=f'Клиент {dfa_number}',
            inn='1234567890',
            insurance_type=insurance_type,
            insurance_period='1 год',
            dfa_number=dfa_number,
            branch=branch,
            manager_name=manager_name,
            deal_status=deal_status,
            response_deadline=response_deadline,
        )
        InsuranceRequest.objects.filter(pk=insurance_request.pk).update(created_at=request_created_at)
        insurance_request.refresh_from_db()

        summary = InsuranceSummary.objects.create(
            request=insurance_request,
            status='completed_accepted',
            selected_company=selected_company,
            selected_franchise_variant=selected_variant,
            completed_at=completed_at,
        )
        InsuranceSummary.objects.filter(pk=summary.pk).update(created_at=summary_created_at)
        summary.refresh_from_db()
        return summary

    def _add_offer(self, summary, company_name, premium, received_at):
        offer = InsuranceOffer.objects.create(
            summary=summary,
            company_name=company_name,
            insurance_year=1,
            insurance_sum=Decimal('1000000.00'),
            premium_with_franchise_1=Decimal(premium),
            franchise_1=Decimal('0'),
            installment_variant_1=True,
            payments_per_year_variant_1=2,
        )
        InsuranceOffer.objects.filter(pk=offer.pk).update(received_at=received_at)

    def test_access_for_admin_and_regular_user(self):
        admin_response = self.client.get(reverse('summaries:analytics_insurance_companies'))
        self.assertEqual(admin_response.status_code, 200)
        self.assertTemplateUsed(admin_response, 'summaries/analytics_insurance_companies.html')

        self.client.logout()
        self.client.login(username='analytics_companies_user', password='testpass123')
        user_response = self.client.get(reverse('summaries:analytics_insurance_companies'))
        self.assertEqual(user_response.status_code, 403)
        self.assertTemplateUsed(user_response, 'insurance_requests/access_denied.html')

    def test_navigation_item_visible_for_admin_only(self):
        analytics_companies_url = reverse('summaries:analytics_insurance_companies')

        admin_response = self.client.get(reverse('summaries:summary_list'))
        self.assertContains(admin_response, analytics_companies_url)

        self.client.logout()
        self.client.login(username='analytics_companies_user', password='testpass123')
        user_response = self.client.get(reverse('summaries:summary_list'))
        self.assertNotContains(user_response, analytics_companies_url)

    def test_page_renders_new_blocks_and_drops_old_ones(self):
        response = self.client.get(reverse('summaries:analytics_insurance_companies'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Куда уходит бизнес')
        self.assertContains(response, 'Доли СК по месяцам')
        self.assertContains(response, 'СК × филиал')
        self.assertContains(response, 'СК × тип страхования')
        self.assertContains(response, 'Качество данных')
        for removed in ('Ценовая конкурентность', 'Конверсия в выбор клиента', 'Менеджер Альянса',
                        'name="date_mode"', 'name="comparison_mode"', 'name="full_coverage"', 'Data Quality'):
            self.assertNotContains(response, removed)

    def test_filters_by_branch_and_type(self):
        url = reverse('summaries:analytics_insurance_companies')

        self.assertEqual(self.client.get(url, {'branch': 'Москва'}).context['kpi']['total_deals'], 1)
        self.assertEqual(
            self.client.get(url, {'insurance_type': 'страхование имущества'}).context['kpi']['total_deals'], 1,
        )

    def test_period_filter_by_summary_creation_date(self):
        url = reverse('summaries:analytics_insurance_companies')
        # Своды созданы 10 и 6 дней назад.
        self.assertEqual(self.client.get(url, {'start_date': (timezone.localdate() - timedelta(days=7)).isoformat()})
                         .context['kpi']['total_deals'], 1)
        self.assertEqual(self.client.get(url, {'period': '180'}).context['kpi']['total_deals'], 2)

    def test_kpi_money_and_cheapest(self):
        response = self.client.get(reverse('summaries:analytics_insurance_companies'))
        kpi = response.context['kpi']

        self.assertEqual(kpi['total_deals'], 2)
        self.assertEqual(kpi['insured_sum_total'], Decimal('2000000.00'))
        # Премии выбранных СК: Абсолют 10 000 + Альфа 8 000.
        self.assertEqual(kpi['premium_total'], Decimal('18000.00'))
        self.assertEqual(kpi['avg_tariff_pct'], Decimal('0.9'))
        self.assertEqual(kpi['winners_count'], 2)
        self.assertEqual(kpi['participants_count'], 3)
        self.assertEqual(kpi['comparable_deals'], 2)
        self.assertEqual(kpi['cheapest_deals'], 1)
        self.assertEqual(kpi['cheapest_rate_pct'], Decimal('50'))

    def test_company_rows(self):
        response = self.client.get(reverse('summaries:analytics_insurance_companies'))
        rows = {row['company_name']: row for row in response.context['company_rows']}

        absolut = rows['Абсолют']
        self.assertEqual(absolut['deals'], 1)
        self.assertEqual(absolut['offered'], 2)
        self.assertEqual(absolut['win_rate_pct'], Decimal('50'))
        self.assertEqual(absolut['premium_total'], Decimal('10000.00'))
        self.assertEqual(absolut['tariff_pct'], Decimal('1'))
        self.assertEqual((absolut['cheapest_wins'], absolut['comparable_wins']), (0, 1))

        alfa = rows['Альфа']
        self.assertEqual((alfa['cheapest_wins'], alfa['comparable_wins']), (1, 1))

        # СК без выигрышей остаётся в таблице — внизу, с участием.
        vsk = rows['ВСК']
        self.assertEqual((vsk['deals'], vsk['offered']), (0, 1))
        self.assertFalse(vsk['is_winner'])
        self.assertEqual(response.context['company_rows'][-1]['company_name'], 'ВСК')

    def test_multiyear_premium_sums_all_years_and_insured_sum_uses_first_year(self):
        InsuranceOffer.objects.create(
            summary=self.summary_moscow, company_name='Абсолют', insurance_year=2,
            insurance_sum=Decimal('800000.00'), premium_with_franchise_1=Decimal('7000.00'),
            franchise_1=Decimal('0'),
        )
        response = self.client.get(reverse('summaries:analytics_insurance_companies'), {'branch': 'Москва'})
        row = response.context['company_rows'][0]

        self.assertEqual(row['insured_sum'], Decimal('1000000.00'))
        self.assertEqual(row['premium_total'], Decimal('17000.00'))
        self.assertEqual(row['tariff_pct'], Decimal('1'))

    def test_heatmaps_and_monthly(self):
        response = self.client.get(reverse('summaries:analytics_insurance_companies'))
        heatmap = response.context['heatmaps']['branch']

        self.assertEqual([row['label'] for row in heatmap['rows']], ['Абсолют', 'Альфа'])
        absolut_cells = {cell['column']: cell['value'] for cell in heatmap['rows'][0]['cells']}
        self.assertEqual(absolut_cells['Москва'], 1)
        self.assertEqual(absolut_cells['Санкт-Петербург'], 0)

        monthly = response.context['charts']['monthly']
        self.assertEqual(sum(sum(series['data']) for series in monthly['deals']), 2)
        # Цвет закреплён за СК, а не за местом в рейтинге.
        self.assertEqual(monthly['colors']['Абсолют'], '#2a78d6')

    def test_data_quality_in_plain_russian(self):
        response = self.client.get(reverse('summaries:analytics_insurance_companies'))
        rows = {row['key']: row for row in response.context['data_quality_rows']}

        self.assertEqual(rows['missing_variant']['count'], 1)
        self.assertNotIn('selected_franchise_variant', rows['missing_variant']['label'])

    def test_export_returns_workbook_with_sheets(self):
        response = self.client.get(reverse('summaries:export_analytics_insurance_companies_widget'), {'period': 'all'})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('analytics_companies_', response['Content-Disposition'])

        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Сводка', 'СК', 'По месяцам', 'СК × филиал', 'СК × тип'])
        companies = workbook['СК']
        self.assertEqual(companies['A5'].value, 'СК')
        self.assertEqual(companies['A6'].value, 'Абсолют')
        self.assertEqual(workbook['Сводка']['B6'].value, 2)
