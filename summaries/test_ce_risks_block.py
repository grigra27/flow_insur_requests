"""Блок «Доп. риски КАСКО C/E» в ответе страховщика V2 (решение владельца 2026-10-14).

Письмо просит расширенные риски для техники C/E, но ≈ 95% ответов о них молчат. Блок обязателен для заявок
C/E с даты выката (CE_RISKS_SINCE); показывается в своде V2 колонкой «Доп. риски C/E»; свод V1 не меняется.
"""
import json
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.templates import EmailTemplateGenerator
from insurance_requests.models import InsuranceRequest
from summaries.models import InsuranceOffer, InsuranceSummary, InsurerResponse
from summaries.response_sections import CE_RISKS, CE_RISKS_SINCE, required_sections
from summaries.response_template import build_v2
from summaries.services.insurer_response import v1_comment_note
from summaries.services.summary_export_v2 import SummaryExportV2Service

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
AFTER = timezone.make_aware(timezone.datetime(CE_RISKS_SINCE.year, CE_RISKS_SINCE.month, CE_RISKS_SINCE.day, 12, 0))
BEFORE = timezone.make_aware(timezone.datetime(2026, 10, 10, 12, 0))


class _Request:
    def __init__(self, has_casco_ce, created_at=None, insurance_type='КАСКО'):
        self.has_casco_ce, self.created_at, self.insurance_type = has_casco_ce, created_at, insurance_type


class CeRisksRuleTests(SimpleTestCase):
    def test_required_only_for_ce_requests_created_since_rollout(self):
        self.assertTrue(CE_RISKS.is_required(_Request(True, AFTER)))
        self.assertTrue(CE_RISKS.is_required(_Request(True, None)))  # новая, ещё не сохранённая заявка
        self.assertFalse(CE_RISKS.is_required(_Request(True, BEFORE)))  # шаблон скачан до выката — без блока
        self.assertFalse(CE_RISKS.is_required(_Request(False, AFTER)))
        self.assertEqual([s.key for s in required_sections(_Request(True, AFTER))], ['inspection', 'ce_risks'])

    def test_choice_labels_have_no_commas_and_aliases_work(self):
        field = CE_RISKS.fields[0]
        for label in field.choice_labels:
            self.assertNotIn(',', label)  # запятая режет пункт списка Excel на два
        self.assertEqual(field.match_choice('включены все'), 'all')
        self.assertEqual(field.match_choice('Частично'), 'partial')
        self.assertEqual(field.match_choice('нет'), 'none')


def _filled(workbook, blocks):
    ws = workbook['Ответ СК']
    ws['B2'], ws['B3'] = 'Согаз', 'Российская Федерация'
    for col, value in zip('BDEF', (7_300_000, 200_000, 0, 1)):
        ws[f'{col}6'] = value
    for name, value in blocks.items():
        sheet, ref = next(iter(workbook.defined_names[name].destinations))
        workbook[sheet][ref.replace('$', '')] = value
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile('ответ.xlsx', buffer.getvalue(), content_type=XLSX)


@override_settings(RESPONSE_TEMPLATE_V2_FOR_ALL=True, RESPONSE_TEMPLATE_V2_SINCE='2026-10-08')
class CeRisksFlowTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.staff = User.objects.create_user('staff', password='p')
        cls.staff.groups.add(Group.objects.get_or_create(name='Пользователи')[0])
        cls.request_obj = InsuranceRequest.objects.create(
            client_name='ООО Тягач', inn='1', dfa_number='ТС-1-ГА-КЗ', object_description='Тягач MAN TGX',
            insurance_type='КАСКО', has_casco_ce=True, condition='new')
        InsuranceRequest.objects.filter(pk=cls.request_obj.pk).update(created_at=AFTER)
        cls.request_obj.refresh_from_db()
        cls.summary = InsuranceSummary.objects.create(request=cls.request_obj, status='collecting')

    def _upload(self, upload):
        self.client.force_login(self.staff)
        response = self.client.post(reverse('summaries:upload_multiple_company_responses', args=[self.summary.pk]),
                                    {'excel_files': [upload]})
        return json.loads(response.content)['results'][0]

    def test_template_has_block_with_dropdown_and_hint(self):
        wb = load_workbook(BytesIO(build_v2(self.request_obj, self.summary.pk)))
        self.assertIn('resp_ce_risks_status', wb.defined_names)
        self.assertIn('resp_ce_risks_comment', wb.defined_names)
        texts = ' '.join(str(c.value) for row in wb['Ответ СК'].iter_rows() for c in row if c.value)
        self.assertIn('Доп. риски КАСКО категории C/E', texts)
        self.assertIn('провал грунта', texts)
        lists = [dv.formula1 for dv in wb['Ответ СК'].data_validations.dataValidation if dv.type == 'list']
        self.assertIn('"Включены все,Включены частично,Не включены"', lists)

    def test_answer_saved_shown_in_summary_v2_not_in_v1(self):
        result = self._upload(_filled(load_workbook(BytesIO(build_v2(self.request_obj, self.summary.pk))), {
            'resp_inspection_status': 'Осмотр не требуется',
            'resp_ce_risks_status': 'Включены частично',
            'resp_ce_risks_comment': 'вне дорог, опрокидывание',
        }))
        self.assertTrue(result['success'], result)
        response = InsurerResponse.objects.get(summary=self.summary, company_name='Согаз')
        self.assertEqual((response.ce_risks_status, response.ce_risks_comment), ('partial', 'вне дорог, опрокидывание'))
        # свод V2: колонка в блоке условий и комментарий «какие риски»
        ws = load_workbook(BytesIO(SummaryExportV2Service().generate(self.summary, is_client_version=True).getvalue()))['Свод']
        titles = [ws.cell(row=9, column=c).value for c in range(1, ws.max_column + 1)]
        col = titles.index('Доп. риски C/E') + 1
        self.assertEqual(ws.cell(row=11, column=col).value, 'Включены частично')
        notes = ws.cell(row=11, column=titles.index('Комментарии') + 1).value
        self.assertIn('Доп. риски C/E: вне дорог, опрокидывание', notes)
        # свод V1 не меняется (решение владельца): в комментарии V1 блок C/E не попадает
        self.assertNotIn('C/E', v1_comment_note(response))

    def test_missing_answer_rejected_for_new_ce_request(self):
        result = self._upload(_filled(load_workbook(BytesIO(build_v2(self.request_obj, self.summary.pk))), {
            'resp_inspection_status': 'Осмотр не требуется',
        }))
        self.assertFalse(result['success'])
        self.assertIn('Доп. риски', result['error_message'])

    def test_old_ce_request_does_not_require_block(self):
        InsuranceRequest.objects.filter(pk=self.request_obj.pk).update(created_at=BEFORE)
        self.request_obj.refresh_from_db()
        self.assertNotIn('ce_risks', [s.key for s in required_sections(self.request_obj)])
        InsuranceOffer.objects.create(summary=self.summary, company_name='Альфа', insurance_year=1,
                                      insurance_sum=Decimal('1'), franchise_1=Decimal('0'),
                                      premium_with_franchise_1=Decimal('1'), coverage_territory='РФ')
        ws = load_workbook(BytesIO(SummaryExportV2Service().generate(self.summary, is_client_version=True).getvalue()))['Свод']
        titles = [ws.cell(row=9, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertNotIn('Доп. риски C/E', titles)

    def test_email_points_to_template_block(self):
        body = EmailTemplateGenerator().generate_email_body(self.request_obj.to_dict() | {'has_casco_ce': True})
        self.assertIn('блок «Доп. риски C/E»', body)
        self.assertNotIn('Просим Вас указывать, что тарифы даны с расширенными рисками.', body)
