"""Шаблон ответа V2, этап 3: тестовая карточка, скачивание шаблона, блоки у компаний на карточке свода.

Всё видно только суперпользователю (контур V2), сотрудники карточку свода видят как раньше.
"""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest
from summaries.models import InsuranceOffer, InsuranceSummary, InsurerResponse
from summaries.response_template import META_SHEET, SHEET_TITLE


class ResponseV2Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        admins = Group.objects.get_or_create(name='Администраторы')[0]
        users = Group.objects.get_or_create(name='Пользователи')[0]
        cls.superuser = User.objects.create_superuser('root', password='p')
        cls.superuser.groups.add(admins)
        cls.admin = User.objects.create_user('admin', password='p')
        cls.admin.groups.add(admins)
        cls.staff = User.objects.create_user('staff', password='p')
        cls.staff.groups.add(users)
        request = InsuranceRequest.objects.create(
            client_name='ООО Тест', inn='1234567890', dfa_number='ОБ-20702-ЛО-КР', vehicle_info='GIANFRANCO',
            insurance_type='страхование имущества', has_transportation=True,
            transportation_departure='Москва', transportation_destination='Армавир',
        )
        cls.summary = InsuranceSummary.objects.create(request=request, status='collecting')
        for company in ('Согаз', 'ВСК'):
            InsuranceOffer.objects.create(
                summary=cls.summary, company_name=company, insurance_year=1,
                insurance_sum=Decimal('3000000'), franchise_1=Decimal('0'),
                premium_with_franchise_1=Decimal('90000'), coverage_territory='РФ',
            )
        InsurerResponse.objects.create(
            summary=cls.summary, company_name='Согаз', template_version=InsurerResponse.TEMPLATE_V2,
            inspection_status='photo', rnpk_status='included', transport_cost=Decimal('5330'),
        )



class ResponseV2UiTests(ResponseV2Base):
    def _detail(self, user):
        self.client.force_login(user)
        return self.client.get(reverse('summaries:summary_detail', args=[self.summary.pk]))

    def test_test_card_visible_only_to_superuser(self):
        self.assertContains(self._detail(self.superuser), 'ТЕСТ · Шаблон ответа V2')
        for user in (self.admin, self.staff):
            response = self._detail(user)
            self.assertEqual(response.status_code, 200)
            self.assertNotContains(response, 'Шаблон ответа V2')

    def test_block_lines_for_superuser_with_missing_marked(self):
        content = self._detail(self.superuser).content.decode()
        self.assertIn('class="og-v2-lines"', content)
        self.assertIn('Требуется осмотр, возможен осмотр по фотографиям', content)
        self.assertIn('Будут прописаны в полисе', content)
        self.assertNotIn('Будут прописаны в полисе ·', content)  # у РНПК нет комментария
        self.assertIn('5 330 ₽', content)
        self.assertIn('og-v2-line og-v2-line--missing', content)  # у ВСК ответа V2 нет
        self.assertNotIn('class="og-v2-lines"', self._detail(self.staff).content.decode())

    def test_personal_template_download(self):
        self.client.force_login(self.superuser)
        response = self.client.get(reverse('summaries:download_response_template_v2', args=[self.summary.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn('otvet_strahovshika_v2_20702.xlsx', response['Content-Disposition'])
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb[META_SHEET]['B2'].value, self.summary.pk)
        self.assertEqual(wb[SHEET_TITLE]['A1'].value, 'Ответ страховой компании на запрос ОБ-20702-ЛО-КР')

    def test_generic_template_download(self):
        self.client.force_login(self.superuser)
        response = self.client.get(reverse('summaries:download_response_template_v2_generic'))
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        self.assertIsNone(wb[META_SHEET]['B2'].value)
        self.assertTrue(wb[SHEET_TITLE]['A12'].value.startswith('Общий шаблон.'))

    def test_downloads_denied_without_superuser(self):
        self.client.force_login(self.admin)
        for url in (reverse('summaries:download_response_template_v2', args=[self.summary.pk]),
                    reverse('summaries:download_response_template_v2_generic')):
            self.assertEqual(self.client.get(url).status_code, 403)


class InsurerResponseManualEditTests(ResponseV2Base):
    """Этап 4: ручная правка блоков «Ответа СК» на карточке свода."""

    def _post(self, user, **data):
        self.client.force_login(user)
        return self.client.post(reverse('summaries:set_insurer_response', args=[self.summary.pk]), data)

    def test_superuser_creates_response_for_company_without_one(self):
        response = self._post(self.superuser, company='ВСК', inspection_status='not_required', rnpk_status='not_included',
                              transport_cost='5 330 руб.', transport_terms='на время перевозки')
        self.assertEqual(response.json(), {'success': True})
        saved = InsurerResponse.objects.get(summary=self.summary, company_name='ВСК')
        self.assertEqual((saved.rnpk_status, saved.inspection_status), ('not_included', 'not_required'))
        self.assertEqual(saved.transport_cost, Decimal('5330.00'))
        self.assertEqual(saved.template_version, InsurerResponse.TEMPLATE_V1)
        self.assertEqual(saved.created_by, self.superuser)

    def test_edit_existing_keeps_version_and_empty_clears(self):
        self._post(self.superuser, company='Согаз', rnpk_status='Не будут прописаны в полисе', transport_cost='')
        saved = InsurerResponse.objects.get(summary=self.summary, company_name='Согаз')
        self.assertEqual(saved.rnpk_status, 'not_included')
        self.assertIsNone(saved.transport_cost)
        self.assertEqual(saved.inspection_status, 'photo')  # поля, которых нет в форме, не трогаются
        self.assertEqual(saved.template_version, InsurerResponse.TEMPLATE_V2)

    def test_invalid_values_rejected_with_message(self):
        for data, expected in (
            ({'rnpk_status': 'может быть'}, 'выберите одно из значений'),
            ({'rnpk_status': 'included', 'transport_cost': 'дорого'}, 'укажите сумму в рублях'),
        ):
            with self.subTest(data=data):
                response = self._post(self.superuser, company='ВСК', **data)
                self.assertEqual(response.status_code, 400)
                self.assertIn(expected, response.json()['error'])
        self.assertFalse(InsurerResponse.objects.filter(company_name='ВСК').exists())

    def test_company_without_offers_rejected(self):
        response = self._post(self.superuser, company='Альфа', rnpk_status='included')
        self.assertEqual(response.status_code, 400)
        self.assertIn('нет предложений', response.json()['error'])

    def test_not_available_outside_v2_contour(self):
        self.assertEqual(self._post(self.staff, company='ВСК', rnpk_status='included').status_code, 403)
        self.assertFalse(InsurerResponse.objects.filter(company_name='ВСК').exists())

    def test_edit_form_on_page_for_superuser_only(self):
        self.client.force_login(self.superuser)
        content = self.client.get(reverse('summaries:summary_detail', args=[self.summary.pk])).content.decode()
        self.assertIn('class="og-v2-form d-none js-response-form" data-company="ВСК"', content)
        self.assertIn('<option value="included" selected>Будут прописаны в полисе</option>', content)  # Согаз
        self.assertIn('name="transport_cost" value="5330"', content)
        self.client.force_login(self.staff)
        content = self.client.get(reverse('summaries:summary_detail', args=[self.summary.pk])).content.decode()
        self.assertNotIn('js-response-form"', content)
        self.assertNotIn('data-response-url', content)


class ManualOfferWithV2BlocksTests(ResponseV2Base):
    """Ручное добавление предложения: блоки V2 в форме (контур V2) и обязательная территория (для всех)."""

    def _post(self, user, company='Пари', territory='Российская Федерация', **extra):
        self.client.force_login(user)
        data = {
            'company_name': company, 'coverage_territory': territory, 'notes': '',
            'payments_per_year_variant_1': '1', 'payments_per_year_variant_2': '1',
            'rows-TOTAL': '2',
            'rows-0-insurance_year': '1', 'rows-0-insurance_sum': '3000000', 'rows-0-franchise_1': '0',
            'rows-0-premium_with_franchise_1': '90000',
            'rows-1-insurance_year': '2', 'rows-1-insurance_sum': '2700000', 'rows-1-franchise_1': '0',
            'rows-1-premium_with_franchise_1': '85000',
        }
        data.update(extra)
        return self.client.post(reverse('summaries:add_offer', args=[self.summary.pk]), data, follow=True)

    def _offers(self, company='Пари'):
        return InsuranceOffer.objects.filter(summary=self.summary, company_name=company)

    def test_superuser_adds_offer_with_blocks(self):
        self._post(self.superuser, inspection_status='Требуется осмотр', rnpk_status='included', transport_cost='5 330')
        self.assertEqual(self._offers().count(), 2)
        response = InsurerResponse.objects.get(summary=self.summary, company_name='Пари')
        self.assertEqual((response.rnpk_status, response.transport_cost), ('included', Decimal('5330.00')))
        self.assertEqual(response.template_version, InsurerResponse.TEMPLATE_V1)

    def test_superuser_missing_required_block_rejected(self):
        page = self._post(self.superuser, inspection_status='required', transport_cost='5330')
        self.assertFalse(self._offers().exists())
        self.assertContains(page, 'Не заполнен блок «Риски РНПК»')
        page = self._post(self.superuser, rnpk_status='included', transport_cost='5330')
        self.assertContains(page, 'Не заполнен блок «Осмотр»')
        self.assertContains(page, 'name="transport_cost" id="v2_transport_cost" value="5330"')  # введённое не теряется

    def test_staff_has_no_v2_section_and_saves_without_blocks(self):
        self.client.force_login(self.staff)
        form_page = self.client.get(reverse('summaries:add_offer', args=[self.summary.pk]))
        self.assertNotContains(form_page, 'Дополнительные условия ответа')
        self._post(self.staff)
        self.assertEqual(self._offers().count(), 2)
        self.assertFalse(InsurerResponse.objects.filter(company_name='Пари').exists())

    def test_form_has_existing_values_for_prefill(self):
        self.client.force_login(self.superuser)
        page = self.client.get(reverse('summaries:add_offer', args=[self.summary.pk]))
        self.assertContains(page, 'Дополнительные условия ответа')
        self.assertContains(page, 'id="response-v2-existing"')
        self.assertContains(page, '"rnpk_status": "included"')  # Согаз

    def test_territory_required_once_for_all_users(self):
        for territory in ('', '—'):
            with self.subTest(territory=territory):
                page = self._post(self.staff, territory=territory)
                self.assertFalse(self._offers().exists())
                self.assertContains(page, 'Укажите территорию страхования', count=2)  # сообщение + поле

    def test_company_territory_edit_rejects_empty(self):
        self.client.force_login(self.staff)
        response = self.client.post(reverse('summaries:set_company_territory', args=[self.summary.pk]),
                                    {'company': 'Согаз', 'territory': '  '})
        self.assertEqual(response.status_code, 400)
        self.assertIn('Укажите территорию страхования', response.json()['error'])
        self.assertTrue(InsuranceOffer.objects.filter(summary=self.summary, company_name='Согаз',
                                                      coverage_territory='РФ').exists())
