"""Территория — одна на страховую компанию в своде (решение владельца 2026-09-27)."""
from decimal import Decimal
from io import BytesIO, StringIO

from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest

from .models import InsuranceOffer, InsuranceSummary
from .services import get_excel_export_service


def make_summary():
    request = InsuranceRequest.objects.create(client_name='ООО Тест', inn='7707083893', dfa_number='ТС-1',
                                              vehicle_info='Haval F7X')
    return InsuranceSummary.objects.create(request=request, status='ready')


def offer(summary, company, year, territory=''):
    return InsuranceOffer.objects.create(summary=summary, company_name=company, insurance_year=year,
                                         insurance_sum=Decimal('3000000'), franchise_1=Decimal('0'),
                                         premium_with_franchise_1=Decimal('70000'), coverage_territory=territory)


def territories(summary, company='ВСК'):
    return list(summary.offers.filter(company_name=company).order_by('insurance_year')
                .values_list('coverage_territory', flat=True))


class TerritorySyncTests(TestCase):
    def setUp(self):
        self.summary = make_summary()

    def test_new_year_without_territory_inherits(self):
        offer(self.summary, 'ВСК', 1, 'РФ и СНГ')
        offer(self.summary, 'ВСК', 2)
        self.assertEqual(territories(self.summary), ['РФ и СНГ', 'РФ и СНГ'])

    def test_editing_one_year_updates_all_years_of_company_only(self):
        first = offer(self.summary, 'ВСК', 1, 'РФ')
        offer(self.summary, 'ВСК', 2, 'РФ')
        offer(self.summary, 'Альфа', 1, 'РФ')
        first.coverage_territory = 'РФ, Европа и СНГ'
        first.save()
        self.assertEqual(territories(self.summary), ['РФ, Европа и СНГ', 'РФ, Европа и СНГ'])
        self.assertEqual(territories(self.summary, 'Альфа'), ['РФ'])

    def test_clearing_territory_clears_all_years(self):
        first = offer(self.summary, 'ВСК', 1, 'РФ')
        offer(self.summary, 'ВСК', 2)
        first.coverage_territory = ''
        first.save()
        self.assertEqual(territories(self.summary), ['', ''])

    def test_saving_other_fields_does_not_touch_territory(self):
        first = offer(self.summary, 'ВСК', 1, 'РФ')
        second = offer(self.summary, 'ВСК', 2)
        InsuranceOffer.objects.filter(pk=second.pk).update(coverage_territory='Особая')  # «старые» данные
        first.premium_with_franchise_1 = Decimal('71000')
        first.save()
        self.assertEqual(territories(self.summary), ['РФ', 'Особая'])

    def test_offer_moved_to_another_company_inherits_there(self):
        offer(self.summary, 'Альфа', 1, 'Вся РФ')
        moved = offer(self.summary, 'ВСК', 2)
        moved.company_name = 'Альфа'
        moved.save()
        self.assertEqual(territories(self.summary, 'Альфа'), ['Вся РФ', 'Вся РФ'])

    def test_set_from_summary_card(self):
        offer(self.summary, 'ВСК', 1, 'РФ')
        offer(self.summary, 'ВСК', 2)
        user = User.objects.create_user(username='terr_user', password='pwd')
        user.groups.add(Group.objects.get_or_create(name='Пользователи')[0])
        self.client.login(username='terr_user', password='pwd')

        url = reverse('summaries:set_company_territory', args=[self.summary.pk])
        response = self.client.post(url, {'company': 'ВСК', 'territory': '  РФ без исключений '})
        self.assertEqual(response.json(), {'success': True, 'updated': 2})
        self.assertEqual(territories(self.summary), ['РФ без исключений', 'РФ без исключений'])
        self.assertEqual(self.client.post(url, {'company': 'Зетта', 'territory': 'x'}).status_code, 400)

        page = self.client.get(reverse('summaries:summary_detail', args=[self.summary.pk]))
        self.assertContains(page, 'js-territory-form')
        self.assertContains(page, 'Территория — для всех лет ВСК')


class TerritoryExcelTests(TestCase):
    def sheet(self, summary):
        data = get_excel_export_service().generate_summary_excel(summary)
        return load_workbook(BytesIO(data.getvalue())).worksheets[0]

    def test_one_text_with_empty_years_is_one_merged_cell(self):
        summary = make_summary()
        offer(summary, 'ВСК', 1, 'РФ, Европа и СНГ')
        second = offer(summary, 'ВСК', 2)
        InsuranceOffer.objects.filter(pk=second.pk).update(coverage_territory='')  # как в старых данных
        ws = self.sheet(summary)
        self.assertIn('J10:J11', {str(rng) for rng in ws.merged_cells.ranges})
        self.assertEqual(ws['J10'].value, 'РФ, Европа и СНГ')

    def test_really_different_texts_stay_per_row(self):
        summary = make_summary()
        offer(summary, 'ВСК', 1, 'РФ')
        second = offer(summary, 'ВСК', 2)
        InsuranceOffer.objects.filter(pk=second.pk).update(coverage_territory='Европа')
        ws = self.sheet(summary)
        self.assertEqual((ws['J10'].value, ws['J11'].value), ('РФ', 'Европа'))


class UnifyCommandTests(TestCase):
    def test_report_then_apply_takes_filled_year(self):
        summary = make_summary()
        offer(summary, 'ВСК', 1, 'РФ, Европа и СНГ')
        second = offer(summary, 'ВСК', 2)
        InsuranceOffer.objects.filter(pk=second.pk).update(coverage_territory='')
        untouched = make_summary()
        offer(untouched, 'Пари', 1)
        offer(untouched, 'Пари', 2)

        out = StringIO()
        call_command('unify_offer_territories', stdout=out)
        self.assertIn('расхождением территории по годам: 1', out.getvalue())
        self.assertEqual(territories(summary), ['РФ, Европа и СНГ', ''])

        call_command('unify_offer_territories', '--apply', stdout=StringIO())
        self.assertEqual(territories(summary), ['РФ, Европа и СНГ', 'РФ, Европа и СНГ'])
        self.assertEqual(territories(untouched, 'Пари'), ['', ''])
