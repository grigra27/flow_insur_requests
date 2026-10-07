"""Шаблон ответа V2, этап 2: загрузка ответа с дополнительными блоками.

docs/improvement_plans/insurer_response_v2.md, §6 и §8. Контур V2 (строгие проверки блоков) —
у суперпользователя, пока RESPONSE_TEMPLATE_V2_FOR_ALL выключен; сотрудники загружают как раньше.
"""
import json
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest
from summaries.exceptions import MissingDataError
from summaries.models import InsuranceCompany, InsuranceOffer, InsuranceSummary, InsurerResponse
from summaries.response_template import SHEET_TITLE, V1_TEMPLATE_PATH, build_v2
from summaries.services.excel_services import ExcelResponseProcessor

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _fill(workbook, company='Согаз', territory='Российская Федерация', blocks=None, years=1):
    ws = workbook[SHEET_TITLE] if SHEET_TITLE in workbook.sheetnames else workbook.active
    ws['B2'], ws['B3'] = company, territory
    for offset in range(years):
        row = 6 + offset
        for col, value in zip('BDEF', (2_000_000 - offset * 100_000, 80_000, 0, 1)):
            ws[f'{col}{row}'] = value
    for name, value in (blocks or {}).items():
        sheet, ref = next(iter(workbook.defined_names[name].destinations))
        workbook[sheet][ref.replace('$', '')] = value
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(content, name='ответ.xlsx'):
    return SimpleUploadedFile(name, content, content_type=XLSX)


class InsurerResponseUploadBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        for order, name in enumerate(['Согаз', 'ВСК'], start=1):
            InsuranceCompany.objects.get_or_create(
                name=name, defaults={'display_name': name, 'sort_order': order, 'is_active': True})
        cls.superuser = User.objects.create_superuser('root', password='p')
        cls.staff = User.objects.create_user('staff', password='p')
        users = Group.objects.get_or_create(name='Пользователи')[0]
        admins = Group.objects.get_or_create(name='Администраторы')[0]
        cls.staff.groups.add(users)
        cls.superuser.groups.add(admins)
        # Имущество с перевозкой (по мотивам ОБ-20702-ЛО-КР): нужны оба блока
        cls.property_request = InsuranceRequest.objects.create(
            client_name='ООО Тест', inn='1234567890', dfa_number='ОБ-20702-ЛО-КР', vehicle_info='GIANFRANCO',
            insurance_type='страхование имущества', has_transportation=True,
            transportation_departure='Москва', transportation_destination='Армавир', transportation_days=3,
        )
        cls.property_summary = InsuranceSummary.objects.create(request=cls.property_request, status='collecting')
        cls.casco_request = InsuranceRequest.objects.create(
            client_name='ООО Тест', inn='1234567890', dfa_number='ТС-20848-ЛА-КЗ', vehicle_info='Mercedes',
            insurance_type='КАСКО',
        )
        cls.casco_summary = InsuranceSummary.objects.create(request=cls.casco_request, status='collecting')

    def _v2(self, summary, blocks=None, generic=False, **kwargs):
        data = build_v2(None if generic else summary.request, None if generic else summary.pk)
        return _fill(load_workbook(BytesIO(data)), blocks=blocks, **kwargs)

    def _v1(self, **kwargs):
        return _fill(load_workbook(V1_TEMPLATE_PATH), **kwargs)

    def _process(self, content, summary, user):
        return ExcelResponseProcessor().process_excel_file(_upload(content), summary, user=user)


class StrictContourTests(InsurerResponseUploadBase):
    """Контур V2 — суперпользователь."""

    FULL_BLOCKS = {'resp_rnpk_status': 'Будут прописаны в полисе', 'resp_rnpk_comment': 'по правилам СК',
                   'resp_transport_cost': 5330, 'resp_transport_terms': 'на время перевозки'}

    def test_filled_v2_file_saves_offers_and_response(self):
        content = self._v2(self.property_summary, self.FULL_BLOCKS, years=2)
        result = self._process(content, self.property_summary, self.superuser)

        self.assertEqual(result['years'], [1, 2])
        self.assertEqual(result['template_version'], 2)
        self.assertEqual(result['response_warnings'], [])
        response = InsurerResponse.objects.get(summary=self.property_summary, company_name='Согаз')
        self.assertEqual(response.template_version, InsurerResponse.TEMPLATE_V2)
        self.assertEqual(response.rnpk_status, 'included')
        self.assertEqual(response.rnpk_comment, 'по правилам СК')
        self.assertEqual(response.transport_cost, Decimal('5330.00'))
        self.assertEqual(response.transport_terms, 'на время перевозки')
        self.assertEqual(response.created_by, self.superuser)
        offer = InsuranceOffer.objects.filter(summary=self.property_summary, company_name='Согаз').first()
        self.assertTrue(response.source_file.name)
        self.assertEqual(response.source_file.name, offer.attachment_file.name)

    def test_empty_required_block_rejects_file(self):
        content = self._v2(self.property_summary, {'resp_transport_cost': 5330})
        with self.assertRaises(MissingDataError) as ctx:
            self._process(content, self.property_summary, self.superuser)
        self.assertIn('не заполнен блок «Риски РНПК». Запросите', str(ctx.exception))
        self.assertIn('Ответ СК «Согаз» не загружен', str(ctx.exception))
        self.assertFalse(InsuranceOffer.objects.filter(summary=self.property_summary).exists())
        self.assertFalse(InsurerResponse.objects.exists())

    def test_dash_in_money_counts_as_empty(self):
        content = self._v2(self.property_summary, {'resp_rnpk_status': 'Не будут прописаны в полисе', 'resp_transport_cost': '—'})
        with self.assertRaises(MissingDataError) as ctx:
            self._process(content, self.property_summary, self.superuser)
        self.assertIn('«Перевозка (транспортировка) предмета лизинга»', str(ctx.exception))

    def test_invalid_choice_and_money_are_rejected(self):
        for blocks, expected in (
            ({'resp_rnpk_status': 'Требуется согласование', 'resp_transport_cost': 100},
             'ожидается одно из: Будут прописаны в полисе, Не будут прописаны в полисе'),
            ({'resp_rnpk_status': 'Будут прописаны в полисе', 'resp_transport_cost': 'дорого'}, 'ожидается сумма в рублях'),
        ):
            with self.subTest(blocks=blocks), self.assertRaises(MissingDataError) as ctx:
                self._process(self._v2(self.property_summary, blocks), self.property_summary, self.superuser)
            self.assertIn(expected, str(ctx.exception))

    def test_hand_written_values_are_normalised(self):
        content = self._v2(self.property_summary,
                           {'resp_rnpk_status': ' риски РНПК не будут прописаны в полисе. ',
                            'resp_transport_cost': '5 330 руб.'})
        self._process(content, self.property_summary, self.superuser)
        response = InsurerResponse.objects.get(summary=self.property_summary)
        self.assertEqual(response.rnpk_status, 'not_included')
        self.assertEqual(response.transport_cost, Decimal('5330.00'))

    def test_old_template_rejected_when_blocks_required(self):
        with self.assertRaises(MissingDataError) as ctx:
            self._process(self._v1(), self.property_summary, self.superuser)
        message = str(ctx.exception)
        self.assertIn('файл в старом шаблоне ответа', message)
        self.assertIn('обязательны блоки «Риски РНПК», «Перевозка (транспортировка) предмета лизинга»', message)

    def test_old_template_accepted_when_no_blocks_required(self):
        result = self._process(self._v1(), self.casco_summary, self.superuser)
        self.assertEqual(result['template_version'], 1)
        response = InsurerResponse.objects.get(summary=self.casco_summary, company_name='Согаз')
        self.assertEqual(response.template_version, InsurerResponse.TEMPLATE_V1)
        self.assertEqual(response.rnpk_status, '')

    def test_generic_template_checked_against_request(self):
        # Общий шаблон: КАСКО — заполненный РНПК не сохраняется (блок по заявке не нужен)
        content = self._v2(self.casco_summary, {'resp_rnpk_status': 'Будут прописаны в полисе'}, generic=True)
        self._process(content, self.casco_summary, self.superuser)
        self.assertEqual(InsurerResponse.objects.get(summary=self.casco_summary).rnpk_status, '')
        # Общий шаблон для имущества с перевозкой — пустые блоки отклоняются, как в персональном
        with self.assertRaises(MissingDataError):
            self._process(self._v2(self.property_summary, generic=True), self.property_summary, self.superuser)

    def test_template_for_other_request_gives_warning(self):
        # шаблон скачан со страницы другой заявки (имущество), загружен в свод КАСКО
        content = _fill(load_workbook(BytesIO(build_v2(self.property_request))))
        result = self._process(content, self.casco_summary, self.superuser)
        self.assertEqual(len(result['response_warnings']), 1)
        self.assertIn(f'для другой заявки (#{self.property_request.pk})', result['response_warnings'][0])
        self.assertIn('ТС-20848-ЛА-КЗ', result['response_warnings'][0])

    def test_early_summary_template_checked_by_summary(self):
        # шаблон, скачанный со свода до появления номера заявки в _meta
        wb = load_workbook(BytesIO(build_v2(self.casco_request, 999)))
        wb['_meta']['A3'] = wb['_meta']['B3'] = None
        result = self._process(_fill(wb), self.casco_summary, self.superuser)
        self.assertIn('сформирован для свода #999', result['response_warnings'][0])

    def test_own_request_template_has_no_warning(self):
        content = _fill(load_workbook(BytesIO(build_v2(self.casco_request))))
        self.assertEqual(self._process(content, self.casco_summary, self.superuser)['response_warnings'], [])

    @override_settings(RESPONSE_TEMPLATE_V2_FOR_ALL=True)
    def test_flag_turns_strict_checks_on_for_staff(self):
        with self.assertRaises(MissingDataError):
            self._process(self._v1(), self.property_summary, self.staff)


class StaffContourTests(InsurerResponseUploadBase):
    """Контур V1 — сотрудники: загрузка как раньше, ничего нового не отклоняется."""

    def test_old_template_works_as_before_without_response_row(self):
        result = self._process(self._v1(), self.property_summary, self.staff)
        self.assertEqual(result['offers_created'], 1)
        self.assertFalse(InsurerResponse.objects.exists())

    def test_v2_file_with_empty_blocks_is_accepted(self):
        result = self._process(self._v2(self.property_summary), self.property_summary, self.staff)
        self.assertEqual(result['offers_created'], 1)
        response = InsurerResponse.objects.get(summary=self.property_summary)
        self.assertEqual((response.rnpk_status, response.transport_cost), ('', None))

    def test_v2_file_values_saved_silently_and_invalid_skipped(self):
        content = self._v2(self.property_summary,
                           {'resp_rnpk_status': 'Включены в полис', 'resp_transport_cost': 'дорого'})  # прежняя подпись
        self._process(content, self.property_summary, self.staff)
        response = InsurerResponse.objects.get(summary=self.property_summary)
        self.assertEqual(response.rnpk_status, 'included')
        self.assertIsNone(response.transport_cost)


class InsurerResponseSyncTests(InsurerResponseUploadBase):
    def test_response_removed_with_last_offer_only(self):
        content = self._v2(self.property_summary, StrictContourTests.FULL_BLOCKS, years=2)
        self._process(content, self.property_summary, self.superuser)
        offers = InsuranceOffer.objects.filter(summary=self.property_summary, company_name='Согаз').order_by('insurance_year')
        offers.first().delete()
        self.assertTrue(InsurerResponse.objects.filter(summary=self.property_summary).exists())
        offers.last().delete()
        self.assertFalse(InsurerResponse.objects.filter(summary=self.property_summary).exists())


class UploadEndpointsTests(InsurerResponseUploadBase):
    def _post(self, user, content):
        self.client.force_login(user)
        response = self.client.post(
            reverse('summaries:upload_multiple_company_responses', args=[self.property_summary.pk]),
            {'excel_files': [_upload(content)]},
        )
        return json.loads(response.content)['results'][0]

    def test_bulk_upload_rejects_for_superuser_and_accepts_for_staff(self):
        rejected = self._post(self.superuser, self._v1())
        self.assertFalse(rejected['success'])
        self.assertIn('файл в старом шаблоне ответа', rejected['error_message'])

        accepted = self._post(self.staff, self._v1())
        self.assertTrue(accepted['success'])
        self.assertEqual(accepted['template_version'], 1)

    def test_bulk_upload_v2_success_reports_version(self):
        result = self._post(self.superuser, self._v2(self.property_summary, StrictContourTests.FULL_BLOCKS))
        self.assertTrue(result['success'], result)
        self.assertEqual(result['template_version'], 2)
        self.assertEqual(result['response_warnings'], [])
