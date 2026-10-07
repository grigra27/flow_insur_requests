"""Шаблон ответа страховщика V2, этап 1: реестр блоков, сборка шаблона, таблица «Ответ СК».

docs/improvement_plans/insurer_response_v2.md. Шаблон V1 должен собираться без изменений;
в V2 базовый блок на прежних адресах, дополнительные блоки — по заявке.
"""
from io import BytesIO

from django.db import IntegrityError, transaction
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest
from summaries.models import InsuranceSummary, InsurerResponse
from summaries.response_sections import (
    RNPK, RNPK_CHOICES, SECTIONS, TRANSPORT, required_sections, transport_route,
)
from summaries.response_template import (
    META_SHEET, SHEET_TITLE, TEMPLATE_VERSION_V2, V1_TEMPLATE_PATH, build_v1, build_v2,
)
from summaries.services.excel_services import ExcelResponseProcessor


def _snapshot(source):
    """Всё, что видит страховщик и читает парсер: значения, стили, объединения, проверки, печать."""
    ws = load_workbook(source).active
    cells = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=10):
        for c in row:
            cells[c.coordinate] = (
                c.value, c.font.name, c.font.sz, c.font.b, c.font.i,
                c.font.color.rgb if c.font.color else None, c.fill.fill_type, c.fill.fgColor.rgb,
                c.alignment.horizontal, c.alignment.vertical, c.alignment.wrap_text, c.border.left.style,
                c.protection.locked, c.number_format,
            )
    return {
        'cells': cells,
        'merged': sorted(map(str, ws.merged_cells.ranges)),
        'validations': sorted((d.type, d.formula1, str(d.sqref)) for d in ws.data_validations.dataValidation),
        'rows': {k: v.height for k, v in ws.row_dimensions.items() if v.height},
        'protection': (ws.protection.sheet, ws.protection.formatRows),
        'print': (ws.page_setup.orientation, ws.print_area),
        'images': len(ws._images),
    }


def _request(**kwargs):
    defaults = dict(client_name='ООО Тест', inn='1234567890', dfa_number='ОБ-20702-ЛО-КР',
                    vehicle_info='GIANFRANCO', insurance_type='КАСКО')
    defaults.update(kwargs)
    return InsuranceRequest(**defaults)


def _name_ref(wb, name):
    """Адрес именованной ячейки: ('Ответ СК', 'D16')."""
    sheet, ref = next(iter(wb.defined_names[name].destinations))
    return sheet, ref.replace('$', '')


class V1TemplateUnchangedTests(SimpleTestCase):
    def test_committed_v1_template_matches_builder(self):
        # Если тест упал после правки раскладки — пересоберите файл scripts/build_flow_answer_template.py
        self.assertEqual(_snapshot(V1_TEMPLATE_PATH), _snapshot(BytesIO(build_v1())))


class ResponseSectionsTests(SimpleTestCase):
    def test_rnpk_for_special_machinery_and_property_only(self):
        self.assertTrue(RNPK.is_required(_request(insurance_type='страхование спецтехники')))
        self.assertTrue(RNPK.is_required(_request(insurance_type='страхование имущества')))
        self.assertFalse(RNPK.is_required(_request(insurance_type='КАСКО')))
        self.assertFalse(RNPK.is_required(None))

    def test_transport_follows_request_flag(self):
        self.assertTrue(TRANSPORT.is_required(_request(has_transportation=True)))
        self.assertFalse(TRANSPORT.is_required(_request(has_transportation=False)))

    def test_required_sections_in_registry_order(self):
        request = _request(insurance_type='страхование имущества', has_transportation=True)
        self.assertEqual(required_sections(request), [RNPK, TRANSPORT])
        self.assertEqual(required_sections(_request()), [])

    def test_transport_route(self):
        request = _request(transportation_departure='Москва', transportation_destination='Армавир',
                           transportation_days=3)
        self.assertEqual(transport_route(request), 'Москва → Армавир · ориентировочно 3 дн.')
        self.assertEqual(transport_route(_request()), '')

    def test_field_names_match_model(self):
        model_fields = {f.name for f in InsurerResponse._meta.get_fields()}
        for section in SECTIONS:
            for fld in section.fields:
                self.assertIn(fld.name, model_fields)


class V2TemplateTests(SimpleTestCase):
    def _load(self, data):
        wb = load_workbook(BytesIO(data))
        return wb, wb[SHEET_TITLE]

    def _block_pill(self, ws, title):
        row = next(r for r in range(12, ws.max_row + 1) if ws[f'A{r}'].value == title)
        return row, ws[f'G{row}'].value

    def test_personal_property_with_transport_requires_both_blocks(self):
        request = _request(insurance_type='страхование имущества', has_transportation=True,
                           transportation_departure='Москва', transportation_destination='Армавир',
                           transportation_days=3)
        wb, ws = self._load(build_v2(request, summary_id=290))
        self.assertEqual(ws['A1'].value, 'Ответ страховой компании на запрос ОБ-20702-ЛО-КР')
        self.assertIn('перевозка: Москва → Армавир · ориентировочно 3 дн.', ws['A12'].value)
        for title in (RNPK.title, TRANSPORT.title):
            _row, pill = self._block_pill(ws, title)
            self.assertEqual(pill, 'ЗАПОЛНИТЕ — обязательно для этого запроса')
        for name in ('resp_rnpk_status', 'resp_rnpk_comment', 'resp_transport_cost', 'resp_transport_terms'):
            sheet, ref = _name_ref(wb, name)
            self.assertEqual(sheet, SHEET_TITLE)
            self.assertFalse(ws[ref].protection.locked, name)
            self.assertIsNone(ws[ref].value, name)
        route_row, _ = self._block_pill(ws, TRANSPORT.title)
        self.assertEqual(ws[f'D{route_row + 1}'].value, 'Москва → Армавир · ориентировочно 3 дн.')
        rnpk_ref = _name_ref(wb, 'resp_rnpk_status')[1]
        lists = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
        self.assertEqual(lists[rnpk_ref], '"' + ','.join(label for _c, label in RNPK_CHOICES) + '"')

    def test_personal_casco_greys_out_blocks(self):
        wb, ws = self._load(build_v2(_request(dfa_number='ТС-20848-ЛА-КЗ'), summary_id=304))
        for title in (RNPK.title, TRANSPORT.title):
            _row, pill = self._block_pill(ws, title)
            self.assertEqual(pill, 'НЕ ТРЕБУЕТСЯ для этого запроса — не заполняйте')
        for name in ('resp_rnpk_status', 'resp_transport_cost'):
            ref = _name_ref(wb, name)[1]
            self.assertTrue(ws[ref].protection.locked, name)
            self.assertEqual(ws[ref].value, '—')
        validated = {str(dv.sqref) for dv in ws.data_validations.dataValidation}
        self.assertNotIn(_name_ref(wb, 'resp_rnpk_status')[1], validated)

    def test_generic_template_shows_conditions_and_keeps_blocks_open(self):
        wb, ws = self._load(build_v2())
        self.assertEqual(ws['A1'].value, 'Ответ страховой компании на запрос котировки')
        self.assertTrue(ws['A12'].value.startswith('Общий шаблон.'))
        for section in SECTIONS:
            _row, pill = self._block_pill(ws, section.title)
            self.assertEqual(pill, section.condition_text)
        for name in ('resp_rnpk_status', 'resp_transport_cost'):
            self.assertFalse(ws[_name_ref(wb, name)[1]].protection.locked, name)

    def test_meta_sheet_is_very_hidden_with_version_and_summary(self):
        wb, _ws = self._load(build_v2(_request(), summary_id=304))
        meta = wb[META_SHEET]
        self.assertEqual(meta.sheet_state, 'veryHidden')
        self.assertEqual((meta['B1'].value, meta['B2'].value), (TEMPLATE_VERSION_V2, 304))
        self.assertIs(wb.active, wb[SHEET_TITLE])
        generic_meta = self._load(build_v2())[0][META_SHEET]
        self.assertIsNone(generic_meta['B2'].value)

    def test_base_block_cells_unchanged(self):
        v1 = _snapshot(BytesIO(build_v1()))['cells']
        v2 = _snapshot(BytesIO(build_v2(_request(insurance_type='страхование спецтехники'))))['cells']
        base = ['B2', 'F2', 'B3'] + [f'{c}{r}' for c in 'ABDEFHIJ' for r in range(4, 11)]
        for coord in base:
            self.assertEqual(v1[coord], v2[coord], coord)


class V2TemplateParsingTests(TestCase):
    def test_current_parser_reads_base_block_of_v2_file(self):
        wb = load_workbook(BytesIO(build_v2(_request(insurance_type='страхование имущества'), summary_id=1)))
        ws = wb[SHEET_TITLE]
        ws['B2'], ws['B3'] = 'Согаз', 'Российская Федерация'
        for col, value in zip('BDEF', (1_000_000, 50_000, 0, 1)):
            ws[f'{col}6'] = value
        data = ExcelResponseProcessor().extract_company_data(ws)
        self.assertEqual(data['company_name'], 'Согаз')
        self.assertEqual([y['year'] for y in data['years']], [1])


class InsurerResponseModelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        request = InsuranceRequest.objects.create(
            client_name='ООО Тест', inn='1234567890', dfa_number='ОБ-1', vehicle_info='Станок',
            insurance_type='страхование имущества',
        )
        cls.summary = InsuranceSummary.objects.create(request=request, status='collecting')

    def test_one_response_per_summary_and_company(self):
        InsurerResponse.objects.create(summary=self.summary, company_name='Согаз', rnpk_status='included')
        with self.assertRaises(IntegrityError), transaction.atomic():
            InsurerResponse.objects.create(summary=self.summary, company_name='Согаз')
        InsurerResponse.objects.create(summary=self.summary, company_name='ВСК')
        self.assertEqual(self.summary.insurer_responses.count(), 2)

    def test_rnpk_choices_come_from_registry(self):
        response = InsurerResponse(summary=self.summary, company_name='Согаз', rnpk_status='approval')
        self.assertEqual(response.get_rnpk_status_display(), 'Требуется согласование')
        self.assertEqual(response.template_version, InsurerResponse.TEMPLATE_V1)
        self.assertEqual(str(response), f'Согаз в своде #{self.summary.pk}')
