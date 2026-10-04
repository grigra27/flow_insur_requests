"""
Шаблон ответа страховщика (templates/flow_answer_template.xlsx) и устойчивость его разбора.

Шаблон собирается scripts/build_flow_answer_template.py; адреса, которые читает
ExcelResponseProcessor, должны оставаться прежними — у страховщиков в обороте
старые копии файла.
"""
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest
from summaries.exceptions import MissingDataError, RowProcessingError
from summaries.models import InsuranceCompany, InsuranceOffer, InsuranceSummary
from summaries.services.excel_services import ExcelResponseProcessor

TEMPLATE_PATH = Path(settings.BASE_DIR) / 'templates' / 'flow_answer_template.xlsx'
INPUT_COLUMNS = 'BDEFHIJ'


def _filled_template(company='Согаз', rows=None, territory='Российская Федерация', notes=None):
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook.active
    sheet['B2'] = company
    sheet['B3'] = territory
    if notes:
        sheet['F2'] = notes
    for offset, values in enumerate(rows or []):
        for column, value in zip(INPUT_COLUMNS, values):
            sheet[f'{column}{6 + offset}'] = value
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class ResponseTemplateStructureTests(TestCase):
    """Структура реального файла шаблона."""

    def setUp(self):
        # Не setUpTestData: Django копирует её атрибуты deepcopy, а лист openpyxl после этого теряет стили
        self.sheet = load_workbook(TEMPLATE_PATH).active
        validation = next(v for v in self.sheet.data_validations.dataValidation if 'B2' in str(v.sqref))
        self.company_names = validation.formula1.strip('"').split(',')

    def test_year_numbers_prefilled_for_all_five_rows(self):
        self.assertEqual([self.sheet[f'A{row}'].value for row in range(6, 11)], [1, 2, 3, 4, 5])

    def test_parser_cells_are_merged_input_areas(self):
        merged = {str(cell_range) for cell_range in self.sheet.merged_cells.ranges}
        self.assertIn('B3:J3', merged)
        self.assertIn('F2:J2', merged)
        self.assertIsNone(self.sheet['B2'].value)
        self.assertIsNone(self.sheet['B3'].value)

    def test_sheet_protected_and_only_input_cells_unlocked(self):
        self.assertTrue(self.sheet.protection.sheet)
        unlocked = ['B2', 'F2', 'B3'] + [f'{c}{r}' for c in INPUT_COLUMNS for r in range(6, 11)]
        for coord in unlocked:
            self.assertFalse(self.sheet[coord].protection.locked, coord)
        for coord in ('A2', 'A6', 'A10', 'B5', 'B13'):
            self.assertTrue(self.sheet[coord].protection.locked, coord)

    def test_validations_cover_company_numbers_and_installments(self):
        by_cell = {}
        for validation in self.sheet.data_validations.dataValidation:
            for cell_range in validation.sqref.ranges:
                by_cell[str(cell_range)] = validation
        self.assertEqual(by_cell['B2'].type, 'list')
        self.assertIn('Совкомбанк СК', by_cell['B2'].formula1)
        self.assertNotIn('Cовкомбанк', by_cell['B2'].formula1)  # латинская «C»
        self.assertIn('Югория', by_cell['B2'].formula1)
        self.assertEqual(by_cell['B6:B10'].type, 'decimal')
        self.assertEqual(by_cell['F6:F10'].type, 'list')
        expected = ','.join(map(str, ExcelResponseProcessor.VALID_INSTALLMENT_VALUES))
        self.assertEqual(by_cell['F6:F10'].formula1, f'"{expected}"')

    def test_company_list_names_pass_matcher_unchanged(self):
        # Справочник как на проде (Югория заведена там через админку)
        for order, name in enumerate(self.company_names, start=1):
            InsuranceCompany.objects.get_or_create(
                name=name, defaults={'display_name': name, 'sort_order': order, 'is_active': True},
            )
        matcher = ExcelResponseProcessor().company_matcher
        for name in self.company_names:
            self.assertEqual(matcher.match_company_name(name), name, name)


class ResponseTemplateImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        for order, name in enumerate(['Согаз', 'Совкомбанк СК', 'Югория'], start=1):
            InsuranceCompany.objects.get_or_create(
                name=name, defaults={'display_name': name, 'sort_order': order, 'is_active': True},
            )
        request = InsuranceRequest.objects.create(
            client_name='ООО Тест', inn='1234567890', dfa_number='ТЕСТ-1', vehicle_info='Автобус',
        )
        cls.summary = InsuranceSummary.objects.create(request=request, status='collecting')

    def _upload(self, content, name='таблица для заполнения.xlsx'):
        return SimpleUploadedFile(
            name, content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def test_filled_template_imports_all_years_and_saves_source_file(self):
        # Реальный ответ Согаза (свод 314 на проде): 4 года, оба варианта
        rows = [
            (17_611_976, 376_897, 0, 2, 360_666, 100_000, 2),
            (14_089_581, 376_897, 0, 2, 360_666, 100_000, 2),
            (11_976_144, 376_897, 0, 2, 360_666, 100_000, 2),
            (10_778_529, 376_894, 0, 2, 360_663, 100_000, 2),
        ]
        content = _filled_template(rows=rows, notes='Требуется осмотр ТС.')

        result = ExcelResponseProcessor().process_excel_file(self._upload(content), self.summary)

        self.assertEqual(result['years'], [1, 2, 3, 4])
        offers = list(InsuranceOffer.objects.filter(summary=self.summary).order_by('insurance_year'))
        self.assertEqual(offers[3].insurance_sum, Decimal('10778529.00'))
        self.assertEqual(offers[0].premium_with_franchise_2, Decimal('360666.00'))
        self.assertEqual(offers[0].payments_per_year_variant_1, 2)
        self.assertEqual(offers[0].notes, 'Требуется осмотр ТС.')
        self.assertTrue(all(o.coverage_territory == 'Российская Федерация' for o in offers))
        # одна копия файла на загрузку, у всех лет одна ссылка
        names = {o.attachment_file.name for o in offers}
        self.assertEqual(len(names), 1)
        stored = offers[0].attachment_file
        self.assertTrue(stored.name.startswith('offers/'))
        with stored.open('rb') as fh:
            self.assertEqual(fh.read(), content)

    def test_empty_template_rejected_for_missing_company(self):
        workbook = load_workbook(BytesIO(_filled_template(company=None, rows=[(1_000_000, 50_000, 0, 1)])))
        with self.assertRaises(MissingDataError) as ctx:
            ExcelResponseProcessor().extract_company_data(workbook.active)
        self.assertIn('B2', str(ctx.exception))

    def test_old_template_placeholder_is_not_silently_other(self):
        workbook = load_workbook(BytesIO(_filled_template(company='Название СК', rows=[(1_000_000, 50_000, 0, 1)])))
        with self.assertRaises(MissingDataError) as ctx:
            ExcelResponseProcessor().extract_company_data(workbook.active)
        self.assertIn('Страховая компания не выбрана', str(ctx.exception))


class TolerantNumberParsingTests(TestCase):
    def setUp(self):
        self.processor = ExcelResponseProcessor()

    def _amount(self, value):
        return self.processor._parse_decimal_with_row(value, 'B6', 'страховая сумма', 6)

    def test_currency_suffixes_and_spaces(self):
        cases = {
            '4 342 000 руб.': Decimal('4342000.00'),
            '4\xa0342\xa0000,50 ₽': Decimal('4342000.50'),
            '138 758,17 р.': Decimal('138758.17'),
            '43420 руб': Decimal('43420.00'),
            '1 000 000 рублей': Decimal('1000000.00'),
            '25000 RUB': Decimal('25000.00'),
        }
        for raw, expected in cases.items():
            self.assertEqual(self._amount(raw), expected, raw)

    def test_float_tail_pasted_as_text_is_decimal_not_thousands(self):
        self.assertEqual(self._amount('247499.99999999996'), Decimal('247500.00'))
        self.assertEqual(self._amount('247499,99999999996'), Decimal('247500.00'))

    def test_three_digit_group_still_thousands(self):
        self.assertEqual(self._amount('1.234'), Decimal('1234.00'))
        self.assertEqual(self._amount('2,000'), Decimal('2000.00'))
        self.assertEqual(self._amount('1.234.567'), Decimal('1234567.00'))

    def test_installment_accepts_integral_decimals(self):
        for raw in (4, 4.0, '4.0', '4,0', ' 12 '):
            self.assertIn(self.processor._parse_installment_with_row(raw, 'F6', 6), (4, 12), raw)

    def test_installment_rejects_fraction_and_text(self):
        for raw in ('2.5', 'ежеквартально', '5'):
            with self.assertRaises(RowProcessingError, msg=raw):
                self.processor._parse_installment_with_row(raw, 'F6', 6)
