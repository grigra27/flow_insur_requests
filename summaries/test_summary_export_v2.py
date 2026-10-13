"""Свод V2 (бета, оформление «Ценники»): содержание, оформление, печать на A4 и доступ только суперпользователю.

Колонки свода КАСКО с двумя предложениями: A — СК, B/C — ИТОГО предложений 1/2, D — год, E — страховая
сумма, F–I и J–M — франшиза, тариф, премия, платежей по предложениям, N — территория, O — осмотр,
P — комментарии. Шапка таблицы — строки 9–10, данные с 11-й."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest
from summaries.models import InsuranceCompany, InsuranceOffer, InsuranceSummary, SummaryCompanyStatus
from summaries.services.summary_export_v2 import SummaryExportV2Service

COMPANIES = ['Альфа', 'ВСК', 'Зетта', 'Согласие', 'Ингосстрах', 'Пари', 'Согаз', 'Ренессанс',
             'Абсолют', 'Росгосстрах', 'Энергогарант', 'Югория']


def _all_text(sheet):
    return ' '.join(str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None)


class SummaryExportV2Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        for order, name in enumerate(COMPANIES, start=1):
            InsuranceCompany.objects.get_or_create(
                name=name, defaults={'display_name': name, 'sort_order': order, 'is_active': True})
        cls.request_obj = InsuranceRequest.objects.create(
            client_name='ООО "ПМК 77"', inn='7751199551', dfa_number='ТС-20848-ЛА-КЗ', branch='Казань',
            insurance_type='КАСКО', brand='Mercedes-Benz', model='V 300 D 4Matic', condition='new',
            acquisition_cost_value=Decimal('17600000'), acquisition_cost_currency='RUB',
            usage_purposes='представительские', anti_theft_systems='Сигнализация: штатная',
            lease_start_date=date(2026, 9, 11), lease_end_date=date(2028, 9, 10),
            franchise_type='both_variants',
        )
        cls.summary = InsuranceSummary.objects.create(request=cls.request_obj, status='collecting')
        # Ответы СК (по мотивам свода 304): ВСК дороже Альфы, у Зетты территория из B3
        cls._offer('ВСК', 1, 17_600_000, 534_670.70, 321_027.57, notes='территория страхования: РФ')
        cls._offer('ВСК', 2, 14_960_000, 534_670.70, 321_027.57)
        cls._offer('Альфа', 1, 17_600_000, 357_773, 285_120)
        cls._offer('Альфа', 2, 14_960_000, 357_773, 285_120)
        cls._offer('Зетта', 1, 17_600_000, 440_000, 352_000, territory='Российская Федерация')
        cls._offer('Зетта', 2, 14_960_000, 440_000, 352_000, territory='Российская Федерация')
        SummaryCompanyStatus.objects.update_or_create(
            summary=cls.summary, company=InsuranceCompany.objects.get(name='Согласие'),
            defaults={'status': SummaryCompanyStatus.DECLINED})

    @classmethod
    def _offer(cls, company, year, insurance_sum, premium_1, premium_2, territory=None, notes='',
               summary=None):
        return InsuranceOffer.objects.create(
            summary=summary or cls.summary, company_name=company, insurance_year=year,
            insurance_sum=Decimal(str(insurance_sum)),
            franchise_1=Decimal('0'), premium_with_franchise_1=Decimal(str(premium_1)),
            franchise_2=Decimal('100000') if premium_2 else None,
            premium_with_franchise_2=Decimal(str(premium_2)) if premium_2 else None,
            coverage_territory=territory, notes=notes,
        )

    def _workbook(self, summary=None, client=True):
        data = SummaryExportV2Service().generate(summary or self.summary, is_client_version=client)
        return load_workbook(BytesIO(data.getvalue()))


class SummaryExportV2ContentTests(SummaryExportV2Base):
    def test_client_has_only_summary_sheet_full_adds_request_parameters(self):
        self.assertEqual(self._workbook(client=True).sheetnames, ['Свод'])
        self.assertEqual(self._workbook(client=False).sheetnames, ['Свод', 'Параметры запроса'])

    def test_header_keeps_v1_fields_and_adds_date(self):
        ws = self._workbook()['Свод']
        self.assertEqual(ws['A1'].value, 'Свод котировок и условий')
        self.assertTrue(ws['A2'].value.startswith('от '))
        labels = [ws[f'A{r}'].value for r in range(4, 9)]
        self.assertEqual(labels, ['Заявка №', 'Объект страхования', 'Лизингополучатель',
                                  'Цели использования', None])
        self.assertEqual(ws['B4'].value, 'ТС-20848-ЛА-КЗ')
        self.assertEqual(ws['B6'].value, 'ООО "ПМК 77"')

    def test_header_shows_note_only_when_filled(self):
        self.summary.notes = 'Клиент просит рассрочку'
        self.summary.save(update_fields=['notes'])
        ws = self._workbook()['Свод']
        self.assertEqual(ws['A8'].value, 'Примечание')
        self.assertEqual(ws['B8'].value, 'Клиент просит рассрочку')

    def test_companies_sorted_by_total_of_offer_1_and_declined_last(self):
        ws = self._workbook()['Свод']
        names = [ws[f'A{r}'].value for r in range(11, ws.max_row + 1) if ws[f'A{r}'].value]
        self.assertEqual(names, ['Альфа', 'Зетта', 'ВСК', 'Согласие'])
        self.assertEqual(ws['B17'].value, 'Отказ от страхования')

    def test_year_rows_rates_and_totals_are_live_formulas(self):
        ws = self._workbook()['Свод']
        self.assertEqual(ws['D11'].value, '1 год')
        self.assertEqual(ws['E11'].value, 17_600_000)
        self.assertEqual(ws['F11'].value, 0)  # франшиза — перед тарифом и премией
        self.assertEqual(ws['G11'].value, '=IF(AND(E11<>0,H11<>0),H11/E11,"")')
        self.assertEqual(ws['H11'].value, 357_773)
        self.assertEqual(ws['J11'].value, 100_000)
        self.assertEqual(ws['B11'].value, '=SUM(H11:H12)')  # ценники — сразу после названия СК
        self.assertEqual(ws['C11'].value, '=SUM(L11:L12)')
        self.assertIn('B11:B12', {str(r) for r in ws.merged_cells.ranges})

    def test_totals_and_rates_have_cached_values_for_protected_view(self):
        # Файл с прода открывается в «защищённом просмотре», где Excel не пересчитывает формулы
        data = SummaryExportV2Service().generate(self.summary, is_client_version=True).getvalue()
        ws = load_workbook(BytesIO(data), data_only=True)['Свод']
        self.assertEqual(ws['B11'].value, 357_773 * 2)  # ИТОГО Альфы
        self.assertEqual(ws['C11'].value, 285_120 * 2)
        self.assertAlmostEqual(ws['G11'].value, 357_773 / 17_600_000)
        self.assertEqual(load_workbook(BytesIO(data))['Свод']['B11'].value, '=SUM(H11:H12)')

    def test_v1_summary_formulas_have_cached_values_too(self):
        from summaries.services import get_excel_export_service

        for client in (True, False):
            data = get_excel_export_service().generate_summary_excel(self.summary, is_client_version=client).getvalue()
            formulas, values = load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)
            cells = [(ws.title, c.coordinate) for ws in formulas for row in ws.iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith('=')]
            self.assertTrue(cells)
            missing = [cell for cell in cells if values[cell[0]][cell[1]].value is None]
            self.assertEqual(missing, [])
            # в V1 формулами остались тарифы (ИТОГО пишется числом); тариф Альфы, предложение 1
            rates = [values[t][c].value for t, c in cells if formulas[t][c].value.startswith('=IF(AND(')]
            self.assertTrue(any(abs(rate - 357_773 / 17_600_000) < 1e-12 for rate in rates))

    def test_variant_captions_and_column_titles(self):
        ws = self._workbook()['Свод']
        self.assertEqual(ws['B9'].value, 'ИТОГО')
        self.assertEqual(ws['B10'].value, 'Предложение 1 · без франшизы')
        self.assertEqual(ws['C10'].value, 'Предложение 2 · франшиза 100 000')
        self.assertEqual(ws['F9'].value, 'Детализация · Предложение 1')
        self.assertEqual(ws['J9'].value, 'Детализация · Предложение 2')
        self.assertEqual([ws[f'{c}10'].value for c in 'FGHI'], ['Франшиза', 'Тариф', 'Премия', 'Платежей в год'])
        self.assertEqual([ws[f'{c}10'].value for c in 'JKLM'], ['Франшиза', 'Тариф', 'Премия', 'Платежей в год'])
        self.assertEqual((ws['D9'].value, ws['E9'].value), ('Год', 'Страховая сумма'))
        self.assertEqual(ws['N9'].value, 'Территория страхования')
        self.assertEqual(ws['O9'].value, 'Осмотр')  # КАСКО: без РНПК и перевозки
        self.assertEqual(ws['P9'].value, 'Комментарии')

    def test_territory_and_notes_follow_v1_rules_without_service_placeholder(self):
        ws = self._workbook()['Свод']
        self.assertEqual(ws['N13'].value, 'Российская Федерация')  # Зетта
        self.assertEqual(ws['N11'].value, '—')  # Альфа: ответ до сбора территории
        self.assertNotIn('ранее не собиралась', _all_text(ws))
        # системный комментарий о франшизе — как в своде V1
        self.assertIn('Требуется согласование франшизы с ГО.', ws['P15'].value)
        self.assertTrue(ws['P15'].value.startswith('территория страхования: РФ'))

    def test_long_text_goes_to_extra_row_and_year_rows_keep_height(self):
        request = InsuranceRequest.objects.create(client_name='ООО Тест', inn='3', dfa_number='ТС-3',
                                                 vehicle_info='Самосвал', franchise_type='both_variants')
        summary = InsuranceSummary.objects.create(request=request, status='collecting')
        long_territory = 'Российская Федерация, за исключением зон вооружённых конфликтов. ' * 6
        for year in (1, 2):
            self._offer('Пари', year, 7_300_000, 205_000, 150_000, territory=long_territory, summary=summary)
        ws = self._workbook(summary)['Свод']
        self.assertEqual((ws['D11'].value, ws['D12'].value, ws['D13'].value), ('1 год', '2 год', None))
        self.assertEqual(ws.row_dimensions[11].height, ws.row_dimensions[12].height)
        self.assertGreater(ws.row_dimensions[13].height, 0)
        merged = {str(r) for r in ws.merged_cells.ranges}
        self.assertIn('N11:N13', merged)  # территория — с добавочной строкой
        self.assertIn('A11:A13', merged)  # СК и ценник — на весь блок, вместе с добавочной строкой
        self.assertIn('B11:B13', merged)
        self.assertEqual(ws['B11'].alignment.vertical, 'center')
        self.assertEqual(ws['F10'].alignment.vertical, 'center')  # шапка таблицы — по центру

    def test_selection_fields_are_manual_yellow_with_lists_and_highlight(self):
        ws = self._workbook()['Свод']
        self.assertEqual(ws['O4'].value, 'Выбрана СК:')  # колонка перед «Комментариями»
        self.assertEqual(ws['O6'].value, 'Выбрано предложение:')
        self.assertIsNone(ws['P4'].value)
        self.assertIsNone(ws['P6'].value)
        self.assertEqual(ws['P4'].fill.fgColor.rgb, '00FFF2B3')
        lists = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
        self.assertEqual(lists['P4'], '"Альфа,Зетта,ВСК"')
        self.assertEqual(lists['P6'], '"Предложение 1,Предложение 2"')
        rules = [rule.formula[0] for rng in ws.conditional_formatting for rule in rng.rules]
        self.assertIn('$P$4="Альфа"', rules)
        self.assertIn('AND($P$4="Альфа",$P$6="Предложение 2")', rules)

    def test_price_tags_style(self):
        ws = self._workbook()['Свод']
        self.assertEqual(ws['A1'].font.color.rgb, '00312E81')  # заголовок — индиго, без баннера
        self.assertEqual(ws['A1'].font.sz, 18)
        self.assertEqual(ws['B10'].fill.fgColor.rgb, '00312E81')  # подпись предложения 1 над ценником
        self.assertEqual(ws['C10'].fill.fgColor.rgb, '000F766E')  # предложение 2 — бирюза
        self.assertEqual(ws['B11'].fill.fgColor.rgb, '00EEF0FF')  # ценник
        self.assertEqual(ws['C11'].fill.fgColor.rgb, '00E3F3F0')
        self.assertEqual(ws['B11'].font.color.rgb, '00312E81')
        self.assertEqual(ws['B11'].font.sz, 16)
        self.assertTrue(ws['B11'].font.b)
        self.assertEqual(ws['A11'].font.sz, 12)
        self.assertEqual(ws['H11'].font.sz, 8.5)  # детализация — мелко и серым
        self.assertEqual(ws['H11'].font.color.rgb, '004B5563')
        self.assertEqual(ws['P4'].fill.fgColor.rgb, '00FFF2B3')  # поля менеджера — жёлтые

    def test_white_label(self):
        wb = self._workbook()
        self.assertFalse(wb.properties.creator)
        text = _all_text(wb['Свод']).lower()
        for brand in ('онлайн', 'брокер', 'insflow'):
            self.assertNotIn(brand, text)
        self.assertFalse(wb['Свод']._images)  # ни своего бренда, ни логотипов страховщиков

    def test_request_parameters_sheet_has_v1_fields_and_new_ones(self):
        ws = self._workbook(client=False)['Параметры запроса']
        values = {ws[f'A{r}'].value: ws[f'B{r}'].value for r in range(1, ws.max_row + 1) if ws[f'A{r}'].value}
        sections = [ws[f'A{r}'].value for r in range(1, ws.max_row + 1)
                    if ws[f'A{r}'].value and ws[f'B{r}'].value is None and r > 1]
        self.assertEqual(sections, ['Заявка', 'Лизингополучатель', 'Предмет лизинга', 'Условия страхования',
                                    'Параметры и риски'])
        self.assertEqual(values['ИНН'], '7751199551')
        self.assertEqual(values['Филиал'], 'Казань')
        self.assertEqual(values['Объект'], 'Mercedes-Benz V 300 D 4Matic')  # марка не дублируется
        self.assertEqual(values['Состояние'], 'Новое')
        self.assertEqual(values['Стоимость'], '17 600 000 руб.')
        self.assertNotIn('Примечание', values)  # стоимость известна — пометки «старый загрузчик» нет
        self.assertEqual(values['Договор лизинга'], '11.09.2026 — 10.09.2028 (2 г.)')
        self.assertIn('Рассрочка', values)  # частота уплаты не распознана — хотя бы да/нет
        self.assertEqual(values['Противоугонные системы'], 'Сигнализация: штатная')
        self.assertEqual(values['Автозапуск'], 'Нет')
        self.assertIn('Запрашиваемые риски', values)
        self.assertNotIn('Менеджер сделки', values)  # white label: нашего сотрудника нет

    def test_request_parameters_sheet_takes_parser_v2_fields(self):
        request = self.request_obj
        request.legal_address = '420000, г. Казань, ул. Тестовая, 1'
        request.postal_address = request.legal_address
        request.business_activity = 'Грузоперевозки'
        request.manager_name = 'Тестова Т. Т.'
        request.manufacturing_year = '2025'
        request.insurance_period = '1 год'
        request.insurance_territory = 'РФ'
        request.premium_frequency = 'quarterly'
        request.has_installment = True
        request.has_transportation = True
        request.transportation_departure, request.transportation_destination = 'Москва', 'Казань'
        request.save()
        ws = self._workbook(client=False)['Параметры запроса']
        values = {ws[f'A{r}'].value: ws[f'B{r}'].value for r in range(1, ws.max_row + 1) if ws[f'A{r}'].value}
        self.assertEqual(values['Юридический адрес'], '420000, г. Казань, ул. Тестовая, 1')
        self.assertEqual(values['Почтовый адрес'], 'совпадает с юридическим')
        self.assertEqual(values['Вид деятельности'], 'Грузоперевозки')
        self.assertEqual(values['Менеджер лизинговой компании'], 'Тестова Т. Т.')
        self.assertEqual(values['Год выпуска'], '2025')
        self.assertEqual(values['Срок страхования'], '1 год')
        self.assertEqual(values['Территория'], 'РФ')
        self.assertIn('рассрочка', values['Уплата премии'])
        self.assertNotIn('Рассрочка', values)
        self.assertEqual(values['Перевозка'], 'Да')
        self.assertIn('Москва', values['Маршрут перевозки'])
        self.assertNotIn('Стоимость перевозки', values)  # просьба к страховщику — только в PDF

    def test_single_variant_summary_has_no_offer_2_columns(self):
        request = InsuranceRequest.objects.create(client_name='ООО Тест', inn='1', dfa_number='ТС-1',
                                                 vehicle_info='Экскаватор')
        summary = InsuranceSummary.objects.create(request=request, status='collecting')
        self._offer('Альфа', 1, 1_000_000, 50_000, None, summary=summary)
        ws = self._workbook(summary)['Свод']
        self.assertEqual([ws[f'{c}10'].value for c in 'EFGH'], ['Франшиза', 'Тариф', 'Премия', 'Платежей в год'])
        self.assertEqual(ws['I9'].value, 'Территория страхования')
        self.assertEqual(ws['J9'].value, 'Осмотр')
        self.assertEqual(ws['K9'].value, 'Комментарии')
        lists = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
        self.assertEqual(lists['K6'], '"Предложение 1"')
        self.assertGreaterEqual(ws.row_dimensions[11].height, 30)  # один год — ценник не обрезается


class SummaryExportV2TableLookTests(SummaryExportV2Base):
    def test_only_landscape_print_setting_and_frozen_panes(self):
        wb = self._workbook(client=False)
        ws = wb['Свод']
        self.assertEqual(ws.freeze_panes, 'D11')  # СК и ценники на экране при прокрутке
        self.assertEqual(ws.page_setup.orientation, 'landscape')
        self.assertIsNone(wb['Параметры запроса'].page_setup.orientation)
        for sheet in wb:
            self.assertFalse(sheet.print_area)
            self.assertFalse(sheet.print_title_rows)
            self.assertFalse(sheet.print_title_cols)
            self.assertEqual(list(sheet.col_breaks.brk), [])
            self.assertEqual(list(sheet.row_breaks.brk), [])
            self.assertIsNone(sheet.page_setup.scale)
            self.assertFalse(sheet.sheet_properties.pageSetUpPr and sheet.sheet_properties.pageSetUpPr.fitToPage)
            self.assertFalse(sheet.oddFooter.center.text)

    def test_grid_separators_between_companies_and_column_groups(self):
        ws = self._workbook()['Свод']
        dark = '00312E81'
        # между СК (Альфа 11–12 | Зетта 13–14) — тёмная линия, между годами одной СК — тонкая светлая
        self.assertEqual((ws['H12'].border.bottom.style, ws['H12'].border.bottom.color.rgb), ('medium', dark))
        self.assertEqual(ws['H13'].border.top.style, 'medium')
        self.assertEqual((ws['H11'].border.bottom.style, ws['H11'].border.bottom.color.rgb), ('thin', '00D5D9E8'))
        # вертикали блоков: перед ИТОГО (B), годом (D), детализацией П1 (F) и П2 (J), территорией (N),
        # ответами СК (O) и комментариями (P); внутри блока — тонкая
        for letter in 'BDFJNOP':
            self.assertEqual(ws[f'{letter}11'].border.left.style, 'medium', letter)
        self.assertEqual(ws['G11'].border.left.style, 'thin')
        self.assertEqual(ws['A11'].border.left.style, 'medium')  # рамка таблицы
        self.assertEqual(ws['A9'].border.top.style, 'medium')

    def test_accent_colors_offer_2_teal_conditions_terracotta(self):
        ws = self._workbook()['Свод']
        self.assertEqual(ws['J9'].font.color.rgb, '000F766E')  # «Детализация · Предложение 2»
        self.assertEqual(ws['F10'].font.color.rgb, '00B4532A')  # «Франшиза» — условие
        self.assertEqual(ws['N9'].fill.fgColor.rgb, '00FBEEE8')  # шапка блока условий
        self.assertEqual(ws['A13'].fill.fgColor.rgb, '00F8F9FD')  # зебра: вторая СК
        self.assertEqual(ws['A17'].fill.fgColor.rgb, '00F6F6F8')  # отказ


class SummaryExportV2AccessTests(SummaryExportV2Base):
    def setUp(self):
        self.superuser = User.objects.create_superuser('root', password='p')
        # страница свода открыта группам; суперпользователи на проде состоят в «Администраторах»
        self.superuser.groups.add(Group.objects.get_or_create(name='Администраторы')[0])
        self.admin = User.objects.create_user('admin', password='p')
        self.admin.groups.add(Group.objects.get_or_create(name='Администраторы')[0])

    def _url(self, version):
        return reverse('summaries:generate_summary_file_v2', args=[self.summary.pk, version])

    def test_superuser_downloads_both_versions_in_any_status(self):
        self.client.force_login(self.superuser)
        for version, sheets in (('client', ['Свод']), ('full', ['Свод', 'Параметры запроса'])):
            response = self.client.get(self._url(version))
            self.assertEqual(response.status_code, 200)
            self.assertIn(f'{version}_svod_v2_20848', response['Content-Disposition'])
            self.assertEqual(load_workbook(BytesIO(response.content)).sheetnames, sheets)

    def test_admin_group_without_superuser_is_denied(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(self._url('client')).status_code, 403)

    def test_unknown_version_is_404(self):
        self.client.force_login(self.superuser)
        self.assertEqual(self.client.get(self._url('pdf')).status_code, 404)

    def test_v2_card_visible_only_to_superuser(self):
        detail = reverse('summaries:summary_detail', args=[self.summary.pk])
        self.client.force_login(self.superuser)
        self.assertContains(self.client.get(detail), 'ТЕСТ · Свод V2')
        self.client.force_login(self.admin)
        response = self.client.get(detail)
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Свод V2')


class SummaryExportV2ResponseColumnsTests(SummaryExportV2Base):
    """Этап 5 плана ответа V2: «Осмотр», «Риски РНПК», «Перевозка, ₽» на второй странице."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        from summaries.models import InsurerResponse
        request = InsuranceRequest.objects.create(
            client_name='ООО Склад', inn='2', dfa_number='ОБ-20702-ЛО-КР', vehicle_info='Компрессор',
            insurance_type='страхование имущества', condition='used', has_transportation=True,
            transportation_departure='Москва', transportation_destination='Армавир',
        )
        cls.property_summary = InsuranceSummary.objects.create(request=request, status='collecting')
        cls._offer('Согаз', 1, 3_000_000, 90_000, None, territory='РФ', summary=cls.property_summary)
        cls._offer('Альфа', 1, 3_000_000, 95_000, None, territory='РФ', summary=cls.property_summary)
        InsurerResponse.objects.create(
            summary=cls.property_summary, company_name='Согаз', rnpk_status='included',
            inspection_status='photo', transport_cost=Decimal('5330'), transport_terms='на время перевозки',
        )

    def _sheet(self):
        return self._workbook(self.property_summary)['Свод']

    def test_columns_order_on_page_two(self):
        ws = self._sheet()
        # один вариант: цифры A..H, дальше территория, осмотр, РНПК, перевозка, комментарии
        self.assertEqual([ws[f'{c}9'].value for c in 'IJKLM'],
                         ['Территория страхования', 'Осмотр', 'Риски РНПК', 'Перевозка, ₽\nоплачивается отдельно',
                          'Комментарии'])
        self.assertEqual(ws['L4'].value, 'Выбрана СК:')
        self.assertEqual(ws['M4'].fill.fgColor.rgb, '00FFF2B3')

    def test_transport_paid_separately_marked_with_footnote(self):
        ws = self._sheet()
        self.assertEqual(ws['B9'].value, 'ИТОГО*')
        footnotes = [ws[f'A{r}'].value for r in range(11, ws.max_row + 1)
                     if str(ws[f'A{r}'].value or '').startswith('* ИТОГО — без перевозки')]
        self.assertEqual(len(footnotes), 1)
        self.assertIn('оплачивается отдельно', footnotes[0])
        footnote_row = next(r for r in range(11, ws.max_row + 1) if ws[f'A{r}'].value == footnotes[0])
        self.assertIn(f'A{footnote_row}:H{footnote_row}', {str(r) for r in ws.merged_cells.ranges})  # под колонками цифр
        casco = self._workbook()['Свод']  # без перевозки — ни звёздочки, ни сноски
        self.assertEqual(casco['B9'].value, 'ИТОГО')
        self.assertNotIn('без перевозки', _all_text(casco))

    def test_values_from_insurer_response_and_missing(self):
        ws = self._sheet()
        rows = {ws[f'A{r}'].value: r for r in range(11, ws.max_row + 1) if ws[f'A{r}'].value}
        sogaz, alfa = rows['Согаз'], rows['Альфа']
        self.assertEqual(ws[f'K{sogaz}'].value, 'Будут прописаны в полисе')
        self.assertEqual(ws[f'L{sogaz}'].value, Decimal('5330.00'))
        self.assertEqual(ws[f'L{sogaz}'].number_format, '#,##0')
        self.assertIn('Условия перевозки: на время перевозки', ws[f'M{sogaz}'].value)
        self.assertEqual((ws[f'K{alfa}'].value, ws[f'L{alfa}'].value), ('нет данных', 'нет данных'))

    def test_inspection_column_is_insurer_answer_not_used_object_rule(self):
        ws = self._sheet()
        rows = {ws[f'A{r}'].value: r for r in range(11, ws.max_row + 1) if ws[f'A{r}'].value}
        # объект б/у, но колонка — только ответ страховщика; системной фразы про осмотр нет
        self.assertEqual(ws[f'J{rows["Согаз"]}'].value, 'Требуется осмотр (возможен по фотографиям)')
        self.assertEqual(ws[f'J{rows["Альфа"]}'].value, 'нет данных')
        for r in rows.values():
            self.assertNotIn('обязателен осмотр', (ws[f'M{r}'].value or '').lower())
        casco = self._workbook()['Свод']
        self.assertEqual(casco['O11'].value, 'нет данных')

