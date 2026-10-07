"""Частичное открытие шаблона ответа V2 сотрудникам (RESPONSE_TEMPLATE_V2_FOR_ALL), решение 2026-10-07.

Сотрудникам — комплект для страховщика, блоки ответа, проверки (только для заявок с даты переключения);
Excel-свод — V1 с ответами блоков в комментариях; свод V2 — по-прежнему только суперпользователю.
"""
import json
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.templates import EmailTemplateGenerator
from insurance_requests.models import InsuranceRequest
from summaries.models import InsuranceOffer, InsuranceSummary, InsurerResponse
from summaries.response_template import V1_TEMPLATE_PATH, build_v2
from summaries.services import get_excel_export_service
from summaries.services.insurer_response import v1_comment_note, v2_strict, v2_visible

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
SINCE = '2026-10-08'


def _filled(workbook, blocks=None):
    ws = workbook['Ответ СК'] if 'Ответ СК' in workbook.sheetnames else workbook.active
    ws['B2'], ws['B3'] = 'Согаз', 'Российская Федерация'
    for col, value in zip('BDEF', (2_000_000, 80_000, 0, 1)):
        ws[f'{col}6'] = value
    for name, value in (blocks or {}).items():
        sheet, ref = next(iter(workbook.defined_names[name].destinations))
        workbook[sheet][ref.replace('$', '')] = value
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile('ответ.xlsx', buffer.getvalue(), content_type=XLSX)


@override_settings(RESPONSE_TEMPLATE_V2_FOR_ALL=True, RESPONSE_TEMPLATE_V2_SINCE=SINCE)
class ResponseTemplateV2RolloutTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        users = Group.objects.get_or_create(name='Пользователи')[0]
        admins = Group.objects.get_or_create(name='Администраторы')[0]
        cls.staff = User.objects.create_user('staff', password='p')
        cls.staff.groups.add(users)
        cls.superuser = User.objects.create_superuser('root', password='p')
        cls.superuser.groups.add(admins)
        cls.new_request = InsuranceRequest.objects.create(
            client_name='ООО Новая', inn='1', dfa_number='ОБ-1', vehicle_info='Станок',
            insurance_type='страхование имущества', condition='used',
        )
        cls.old_request = InsuranceRequest.objects.create(
            client_name='ООО Старая', inn='2', dfa_number='ОБ-2', vehicle_info='Станок',
            insurance_type='страхование имущества', condition='used',
        )
        InsuranceRequest.objects.filter(pk=cls.new_request.pk).update(
            created_at=timezone.make_aware(timezone.datetime(2026, 10, 9, 12, 0)))
        InsuranceRequest.objects.filter(pk=cls.old_request.pk).update(
            created_at=timezone.make_aware(timezone.datetime(2026, 10, 1, 12, 0)))
        cls.new_request.refresh_from_db()
        cls.old_request.refresh_from_db()
        cls.new_summary = InsuranceSummary.objects.create(request=cls.new_request, status='collecting')
        cls.old_summary = InsuranceSummary.objects.create(request=cls.old_request, status='collecting')

    def _upload(self, user, summary, upload):
        self.client.force_login(user)
        response = self.client.post(
            reverse('summaries:upload_multiple_company_responses', args=[summary.pk]), {'excel_files': [upload]})
        return json.loads(response.content)['results'][0]

    def test_visibility_and_strictness_rules(self):
        self.assertTrue(v2_visible(self.staff))
        self.assertTrue(v2_strict(self.staff, self.new_request))
        self.assertFalse(v2_strict(self.staff, self.old_request))
        self.assertTrue(v2_strict(self.superuser, self.old_request))
        with override_settings(RESPONSE_TEMPLATE_V2_SINCE=''):
            self.assertTrue(v2_strict(self.staff, self.old_request))

    def test_staff_gets_kit_without_test_tags(self):
        self.client.force_login(self.staff)
        page = self.client.get(reverse('insurance_requests:request_detail', args=[self.new_request.pk]))
        self.assertContains(page, 'Отправка страховщику')
        self.assertNotContains(page, '<span class="rq-kit__tag">')
        self.assertContains(page, 'Заявка для страховщика (PDF)')
        self.assertContains(page, 'Шаблон ответа (Excel)')
        self.assertNotContains(page, 'class="rq-pdf-btn"')
        download = self.client.get(reverse('insurance_requests:download_response_template',
                                           args=[self.new_request.pk]))
        self.assertEqual(download.status_code, 200)

    def test_staff_summary_page_lines_generic_template_and_no_summary_v2(self):
        InsuranceOffer.objects.create(summary=self.new_summary, company_name='Согаз', insurance_year=1,
                                      insurance_sum=Decimal('1000000'), franchise_1=Decimal('0'),
                                      premium_with_franchise_1=Decimal('50000'), coverage_territory='РФ')
        self.client.force_login(self.staff)
        page = self.client.get(reverse('summaries:summary_detail', args=[self.new_summary.pk])).content.decode()
        self.assertIn('class="og-v2-lines"', page)
        self.assertNotIn('<span class="og-v2-tag">V2</span>', page)
        self.assertIn('Общий шаблон ответа (резерв)', page)
        self.assertNotIn('Свод V2', page)
        url = reverse('summaries:generate_summary_file_v2', args=[self.new_summary.pk, 'client'])
        self.assertEqual(self.client.get(url).status_code, 403)

    def test_new_request_rejects_old_template_old_request_accepts(self):
        rejected = self._upload(self.staff, self.new_summary, _filled(load_workbook(V1_TEMPLATE_PATH)))
        self.assertFalse(rejected['success'])
        self.assertIn('файл в старом шаблоне ответа', rejected['error_message'])
        accepted = self._upload(self.staff, self.old_summary, _filled(load_workbook(V1_TEMPLATE_PATH)))
        self.assertTrue(accepted['success'], accepted)

    def test_new_request_v2_file_saved_and_goes_to_v1_comments(self):
        upload = _filled(load_workbook(BytesIO(build_v2(self.new_request, self.new_summary.pk))), {
            'resp_inspection_status': 'Осмотр не требуется',
            'resp_rnpk_status': 'Не будут прописаны в полисе',
        })
        result = self._upload(self.staff, self.new_summary, upload)
        self.assertTrue(result['success'], result)
        response = InsurerResponse.objects.get(summary=self.new_summary, company_name='Согаз')
        self.assertEqual(v1_comment_note(response),
                         'Осмотр: осмотр не требуется. Риски РНПК не будут прописаны в полисе.')
        ws = load_workbook(BytesIO(get_excel_export_service().generate_summary_excel(
            self.new_summary, is_client_version=True).getvalue())).worksheets[0]
        notes_col = next(c.column_letter for c in ws[8] if c.value == 'Комментарии')
        comment = ws[f'{notes_col}10'].value
        self.assertTrue(comment.startswith('Осмотр: осмотр не требуется. Риски РНПК не будут прописаны в полисе.'))
        # страховщик ответил про осмотр — системная фраза для б/у не добавляется
        self.assertNotIn('Обязателен осмотр предмета лизинга', comment)

    def test_v1_system_inspection_note_kept_without_answer(self):
        InsuranceOffer.objects.create(summary=self.old_summary, company_name='ВСК', insurance_year=1,
                                      insurance_sum=Decimal('1000000'), franchise_1=Decimal('0'),
                                      premium_with_franchise_1=Decimal('50000'), coverage_territory='РФ',
                                      notes='без ОСВР')
        ws = load_workbook(BytesIO(get_excel_export_service().generate_summary_excel(
            self.old_summary, is_client_version=True).getvalue())).worksheets[0]
        notes_col = next(c.column_letter for c in ws[8] if c.value == 'Комментарии')
        self.assertEqual(ws[f'{notes_col}10'].value, 'без ОСВР Обязателен осмотр предмета лизинга.')

    def test_add_form_blocks_optional_for_old_request(self):
        self.client.force_login(self.staff)
        data = {'company_name': 'Пари', 'coverage_territory': 'РФ', 'payments_per_year_variant_1': '1',
                'payments_per_year_variant_2': '1', 'rows-TOTAL': '1', 'rows-0-insurance_year': '1',
                'rows-0-insurance_sum': '1000000', 'rows-0-franchise_1': '0',
                'rows-0-premium_with_franchise_1': '50000'}
        self.client.post(reverse('summaries:add_offer', args=[self.old_summary.pk]), data)
        self.assertTrue(InsuranceOffer.objects.filter(summary=self.old_summary, company_name='Пари').exists())
        self.assertFalse(InsurerResponse.objects.filter(summary=self.old_summary, company_name='Пари').exists())
        page = self.client.post(reverse('summaries:add_offer', args=[self.new_summary.pk]), data, follow=True)
        self.assertFalse(InsuranceOffer.objects.filter(summary=self.new_summary, company_name='Пари').exists())
        self.assertContains(page, 'Не заполнен блок «Осмотр»')

    def test_help_and_email_text(self):
        self.client.force_login(self.staff)
        self.assertContains(self.client.get(reverse('summaries:help')), 'Шаблон ответа — свой для каждой заявки')
        body = EmailTemplateGenerator().generate_email_body({
            'insurance_type': 'страхование имущества', 'inn': '1', 'franchise_type': 'none',
            'has_installment': False, 'insurance_period': '1 год',
        })
        self.assertIn('в приложенном шаблоне ответа (блок «Риски РНПК»)', body)
        self.assertIn('необходимость осмотра предмета лизинга', body)
        self.assertNotIn('прилагаемую таблицу', body)


class ResponseTemplateV2FlagOffTests(TestCase):
    def test_email_unchanged_and_staff_not_visible(self):
        staff = User.objects.create_user('s2', password='p')
        self.assertFalse(v2_visible(staff))
        body = EmailTemplateGenerator().generate_email_body({
            'insurance_type': 'страхование имущества', 'inn': '1', 'franchise_type': 'none',
            'has_installment': False, 'insurance_period': '1 год',
        })
        self.assertIn('1. Будут ли включены в полис риски РНПК', body)
        self.assertIn('прилагаемую таблицу', body)
