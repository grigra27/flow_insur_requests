from decimal import Decimal
from io import BytesIO
import importlib
import json
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest
from summaries.forms import AddOfferToSummaryForm, OfferForm
from summaries.models import InsuranceCompany, InsuranceOffer, InsuranceSummary
from summaries.services.excel_services import ExcelExportService, ExcelResponseProcessor


class MockCell:
    def __init__(self, value=None):
        self.value = value


class MockWorksheet:
    def __init__(self, values):
        self.values = values

    def __getitem__(self, cell_address):
        return MockCell(self.values.get(cell_address))


class OfferCoverageTerritoryImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        InsuranceCompany.objects.get_or_create(
            name='Абсолют',
            defaults={
                'display_name': 'Абсолют',
                'sort_order': 1,
                'is_active': True,
            },
        )
        request = InsuranceRequest.objects.create(
            client_name='ООО Тест',
            inn='1234567890',
            dfa_number='ТЕСТ-1',
            vehicle_info='Тестовый объект',
        )
        cls.summary = InsuranceSummary.objects.create(request=request, status='collecting')

    def _worksheet(self, territory=None):
        return MockWorksheet({
            'B2': 'Абсолют',
            'B3': territory,
            'A6': 1,
            'B6': 1_000_000,
            'D6': 50_000,
            'E6': 0,
            'F6': 1,
            'A7': 2,
            'B7': 900_000,
            'D7': 45_000,
            'E7': 0,
            'F7': 1,
        })

    def test_import_copies_common_territory_to_every_year(self):
        processor = ExcelResponseProcessor()
        data = processor.extract_company_data(
            self._worksheet('Российская Федерация, кроме новых территорий')
        )

        offers = processor.create_offers(data, self.summary)

        self.assertEqual(len(offers), 2)
        self.assertEqual(
            {offer.coverage_territory for offer in offers},
            {'Российская Федерация, кроме новых территорий'},
        )

    def test_old_response_without_territory_remains_supported(self):
        processor = ExcelResponseProcessor()

        data = processor.extract_company_data(self._worksheet())

        self.assertEqual(data['coverage_territory'], '')
        offers = processor.create_offers(data, self.summary)
        self.assertTrue(all(offer.coverage_territory == '' for offer in offers))


class OfferCoverageTerritoryExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        request = InsuranceRequest.objects.create(
            client_name='ООО Экспорт',
            inn='1234567890',
            dfa_number='ТЕСТ-2',
            vehicle_info='Складской комплекс',
            insurance_type='страхование имущества',
            insurance_territory='Российская Федерация',
        )
        cls.summary = InsuranceSummary.objects.create(request=request, status='ready')
        cls.offer1 = InsuranceOffer.objects.create(
            summary=cls.summary,
            company_name='Абсолют',
            insurance_year=1,
            insurance_sum=Decimal('1000000'),
            franchise_1=Decimal('0'),
            premium_with_franchise_1=Decimal('50000'),
            franchise_2=Decimal('50000'),
            premium_with_franchise_2=Decimal('45000'),
            coverage_territory='Российская Федерация',
        )
        cls.offer2 = InsuranceOffer.objects.create(
            summary=cls.summary,
            company_name='Абсолют',
            insurance_year=2,
            insurance_sum=Decimal('900000'),
            franchise_1=Decimal('0'),
            premium_with_franchise_1=Decimal('45000'),
            franchise_2=Decimal('50000'),
            premium_with_franchise_2=Decimal('40000'),
            coverage_territory='Российская Федерация',
        )

    def setUp(self):
        self.service = ExcelExportService(
            Path(settings.BASE_DIR) / 'templates' / 'summary_template.xlsx'
        )

    def _generate(self, is_client_version=False):
        return load_workbook(
            self.service.generate_summary_excel(
                self.summary,
                is_client_version=is_client_version,
            )
        )

    def test_all_four_exports_contain_coverage_territory(self):
        for is_client_version in (False, True):
            workbook = self._generate(is_client_version)
            sheet = workbook['summary_template_sheet']
            self.assertEqual(sheet['P10'].value, 'Российская Федерация')
            self.assertIn('P10:P11', {str(cell_range) for cell_range in sheet.merged_cells.ranges})
            if not is_client_version:
                self.assertEqual(
                    workbook['tech_info']['A27'].value,
                    'Территория по исходной заявке:',
                )
                self.assertEqual(workbook['tech_info']['B27'].value, 'Российская Федерация')

        InsuranceOffer.objects.filter(summary=self.summary).update(
            premium_with_franchise_2=None,
            franchise_2=None,
        )
        for is_client_version in (False, True):
            workbook = self._generate(is_client_version)
            sheet = workbook['summary_template_sheet']
            self.assertEqual(sheet['J10'].value, 'Российская Федерация')
            self.assertIn('J10:J11', {str(cell_range) for cell_range in sheet.merged_cells.ranges})

    def test_different_territories_are_not_merged(self):
        self.offer2.coverage_territory = 'Российская Федерация и Республика Беларусь'
        self.offer2.save(update_fields=['coverage_territory'])

        sheet = self._generate()['summary_template_sheet']

        self.assertNotIn('P10:P11', {str(cell_range) for cell_range in sheet.merged_cells.ranges})
        self.assertEqual(sheet['P10'].value, 'Российская Федерация')
        self.assertEqual(
            sheet['P11'].value,
            'Российская Федерация и Республика Беларусь',
        )

    def test_missing_territory_is_explicit_in_export(self):
        InsuranceOffer.objects.filter(summary=self.summary).update(coverage_territory='')

        sheet = self._generate()['summary_template_sheet']

        self.assertEqual(sheet['P10'].value, 'Не указано страховщиком')
        self.assertIn('P10:P11', {str(cell_range) for cell_range in sheet.merged_cells.ranges})

    def test_legacy_territory_is_distinct_from_explicitly_missing(self):
        self.offer1.coverage_territory = None
        self.offer1.save(update_fields=['coverage_territory'])
        self.offer2.coverage_territory = ''
        self.offer2.save(update_fields=['coverage_territory'])

        sheet = self._generate()['summary_template_sheet']

        self.assertEqual(sheet['P10'].value, 'Нет данных: территория ранее не собиралась')
        self.assertEqual(sheet['P11'].value, 'Не указано страховщиком')
        self.assertNotIn('P10:P11', {str(cell_range) for cell_range in sheet.merged_cells.ranges})

    def test_export_sanitizes_illegal_characters_and_truncates_text(self):
        self.offer2.delete()
        self.offer1.coverage_territory = (
            'Россия\x03\x08\x0b\x0e\x1f\r\n' + 'А' * 32768
        )
        self.offer1.notes = 'Комментарий\x03\x08\x0b\x0e\x1f ' + 'Б' * 1100
        self.offer1.save(update_fields=['coverage_territory', 'notes'])

        sheet = self._generate()['summary_template_sheet']

        territory = sheet['P10'].value
        notes = sheet['Q10'].value
        self.assertTrue(territory.startswith('Россия\n'))
        self.assertEqual(len(territory), 32767)
        self.assertTrue(territory.endswith('...'))
        self.assertNotRegex(territory, r'[\x00-\x08\x0b-\x0c\x0e-\x1f]')
        self.assertEqual(len(notes), self.service.MAX_NOTES_LENGTH)
        self.assertTrue(notes.endswith('...'))
        self.assertNotRegex(notes, r'[\x00-\x08\x0b-\x0c\x0e-\x1f]')


class OfferCoverageTerritoryInterfaceAndTemplateTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        group = Group.objects.create(name='Пользователи')
        cls.user = User.objects.create_user(username='territory-user', password='testpass')
        cls.user.groups.add(group)
        request = InsuranceRequest.objects.create(
            client_name='ООО Интерфейс',
            inn='1234567890',
            dfa_number='ТЕСТ-3',
            vehicle_info='Тестовый объект',
            created_by=cls.user,
        )
        cls.summary = InsuranceSummary.objects.create(request=request, status='collecting')
        cls.offer = InsuranceOffer.objects.create(
            summary=cls.summary,
            company_name='Абсолют',
            insurance_year=1,
            insurance_sum=Decimal('1000000'),
            franchise_1=Decimal('0'),
            premium_with_franchise_1=Decimal('50000'),
        )

    def test_forms_expose_optional_coverage_territory(self):
        for form in (OfferForm(), AddOfferToSummaryForm()):
            self.assertIn('coverage_territory', form.fields)
            self.assertFalse(form.fields['coverage_territory'].required)

    def test_summary_page_warns_when_territory_is_missing(self):
        self.offer.coverage_territory = None
        self.offer.save(update_fields=['coverage_territory'])
        self.client.force_login(self.user)

        response = self.client.get(
            reverse('summaries:summary_detail', kwargs={'pk': self.summary.pk})
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Нет данных: предложение создано до сбора территории')
        self.assertContains(response, 'Скачать актуальный шаблон ответа')

    def test_summary_page_distinguishes_explicitly_missing_territory(self):
        self.offer.coverage_territory = ''
        self.offer.save(update_fields=['coverage_territory'])
        self.client.force_login(self.user)

        response = self.client.get(
            reverse('summaries:summary_detail', kwargs={'pk': self.summary.pk})
        )

        self.assertContains(response, 'Не указано страховщиком')

    def test_data_migration_marks_old_blank_values_as_legacy_only(self):
        from django.apps import apps as django_apps

        migration = importlib.import_module(
            'summaries.migrations.0020_insuranceoffer_legacy_coverage_territory'
        )
        migration.mark_legacy_territory_as_unknown(django_apps, None)
        self.offer.refresh_from_db()
        self.assertIsNone(self.offer.coverage_territory)

        new_offer = InsuranceOffer.objects.create(
            summary=self.summary,
            company_name='Абсолют',
            insurance_year=6,
            insurance_sum=Decimal('1000000'),
            franchise_1=Decimal('0'),
            premium_with_franchise_1=Decimal('50000'),
        )
        self.assertEqual(new_offer.coverage_territory, '')

    def test_current_response_template_can_be_downloaded(self):
        self.client.force_login(self.user)

        response = self.client.get(
            reverse('summaries:download_company_response_template')
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], (
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        ))
        self.assertIn('flow_answer_template.xlsx', response['Content-Disposition'])
        workbook = load_workbook(BytesIO(b''.join(response.streaming_content)))
        sheet = workbook['1']
        self.assertEqual(sheet['A3'].value, 'Территория страхования')
        self.assertIn('B3:C3', {str(cell_range) for cell_range in sheet.merged_cells.ranges})

    def test_response_template_download_requires_authentication(self):
        response = self.client.get(
            reverse('summaries:download_company_response_template')
        )

        self.assertEqual(response.status_code, 302)

    def _offer_form_data(self, insurance_year, territory):
        return {
            'company_name': 'Абсолют',
            'insurance_year': insurance_year,
            'insurance_sum': '1000000',
            'franchise_1': '0',
            'premium_with_franchise_1': '50000',
            'franchise_2': '',
            'premium_with_franchise_2': '',
            'payments_per_year_variant_1': '1',
            'payments_per_year_variant_2': '1',
            'coverage_territory': territory,
            'notes': '',
        }

    def _response_file(self, year, territory):
        workbook = load_workbook(
            Path(settings.BASE_DIR) / 'templates' / 'flow_answer_template.xlsx'
        )
        sheet = workbook['1']
        sheet['B2'] = 'Абсолют'
        sheet['B3'] = territory
        sheet['A6'] = year
        sheet['B6'] = 1_000_000
        sheet['D6'] = 50_000
        sheet['E6'] = 0
        sheet['F6'] = 1
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile(
            f'absolute-{year}.xlsx',
            output.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            ),
        )

    def test_manual_add_edit_and_copy_preserve_territory(self):
        self.client.force_login(self.user)

        add_response = self.client.post(
            reverse('summaries:add_offer', args=[self.summary.pk]),
            self._offer_form_data(2, 'Россия и Беларусь'),
        )
        self.assertEqual(add_response.status_code, 302)
        added_offer = InsuranceOffer.objects.get(
            summary=self.summary,
            company_name='Абсолют',
            insurance_year=2,
        )
        self.assertEqual(added_offer.coverage_territory, 'Россия и Беларусь')

        edit_response = self.client.post(
            reverse('summaries:edit_offer', args=[added_offer.pk]),
            self._offer_form_data(2, 'Российская Федерация'),
        )
        self.assertEqual(edit_response.status_code, 302)
        added_offer.refresh_from_db()
        self.assertEqual(added_offer.coverage_territory, 'Российская Федерация')

        copy_response = self.client.post(
            reverse('summaries:copy_offer', args=[added_offer.pk]),
            self._offer_form_data(3, 'Российская Федерация, кроме Крыма'),
        )
        self.assertEqual(copy_response.status_code, 302)
        copied_offer = InsuranceOffer.objects.get(
            summary=self.summary,
            company_name='Абсолют',
            insurance_year=3,
        )
        self.assertEqual(
            copied_offer.coverage_territory,
            'Российская Федерация, кроме Крыма',
        )

    def test_multiple_upload_endpoint_returns_territory_metadata(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse(
                'summaries:upload_multiple_company_responses',
                args=[self.summary.pk],
            ),
            {'excel_files': [self._response_file(4, 'Российская Федерация')]},
        )

        self.assertEqual(response.status_code, 200)
        result = json.loads(response.content)['results'][0]
        self.assertTrue(result['success'])
        self.assertEqual(result['coverage_territory'], 'Российская Федерация')
        self.assertFalse(result['coverage_territory_missing'])
        self.assertEqual(
            InsuranceOffer.objects.get(
                summary=self.summary,
                company_name='Абсолют',
                insurance_year=4,
            ).coverage_territory,
            'Российская Федерация',
        )

        missing_response = self.client.post(
            reverse(
                'summaries:upload_multiple_company_responses',
                args=[self.summary.pk],
            ),
            {'excel_files': [self._response_file(5, None)]},
        )
        missing_result = json.loads(missing_response.content)['results'][0]
        self.assertTrue(missing_result['success'])
        self.assertTrue(missing_result['coverage_territory_missing'])
        self.assertEqual(missing_result['coverage_territory'], '')

    def test_real_templates_have_new_cells_and_safe_rate_formulas(self):
        templates_dir = Path(settings.BASE_DIR) / 'templates'
        response_template = load_workbook(templates_dir / 'flow_answer_template.xlsx')
        response_sheet = response_template['1']
        self.assertEqual(response_sheet['A3'].value, 'Территория страхования')
        self.assertIn('B3:C3', {str(cell_range) for cell_range in response_sheet.merged_cells.ranges})
        self.assertIsNone(response_sheet['B3'].value)

        expected = (
            ('summary_template.xlsx', 'P'),
            ('client_summary_template.xlsx', 'P'),
            ('summary_template_simplified.xlsx', 'J'),
            ('client_summary_template_simplified.xlsx', 'J'),
        )
        for filename, territory_column in expected:
            workbook = load_workbook(templates_dir / filename, data_only=False)
            sheet = workbook['summary_template_sheet']
            self.assertEqual(
                sheet[f'{territory_column}8'].value,
                'Территория страхования',
            )
            self.assertEqual(
                sheet['E10'].value,
                '=IF(AND(C10<>0,F10<>0),F10/C10,"")',
            )
