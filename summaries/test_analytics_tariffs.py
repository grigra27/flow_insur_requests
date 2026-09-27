"""Аналитика тарифов (tariffs_analytics_2026_09, шаг 2)."""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from insurance_requests.models import InsuranceRequest

from .models import InsuranceOffer, InsuranceSummary
from .services import analytics_tariffs as tariffs


def make(text, offers, insurance_type='КАСКО', winner='', condition='new'):
    """offers: {СК: (премия-1, франшиза-1)} или {СК: премия-1}; страховая сумма 1 000 000."""
    request = InsuranceRequest.objects.create(client_name='ООО Тест', inn='7707083893', insurance_type=insurance_type,
                                              vehicle_info=text, dfa_number='ТС-1', condition=condition)
    summary = InsuranceSummary.objects.create(request=request, status='completed_accepted' if winner else 'sent',
                                              selected_company=winner or None)
    for company, value in offers.items():
        premium, franchise = value if isinstance(value, tuple) else (value, 0)
        InsuranceOffer.objects.create(summary=summary, company_name=company, insurance_year=1,
                                      insurance_sum=Decimal('1000000'), franchise_1=Decimal(franchise),
                                      premium_with_franchise_1=Decimal(premium))
    return summary


class TariffServiceTests(TestCase):
    def setUp(self):
        for index in range(3):
            make('Автомобиль Haval M6', {'Зетта': 20000, 'Ингосстрах': 40000, 'Альфа': 30000},
                 winner='Зетта' if index == 0 else '')
        make('Gazelle NEXT', {'Альфа': 25000, 'Согаз': 22000})
        make('Экскаватор SANY SY75C', {'Абсолют': 3000, 'Согаз': 6000}, insurance_type='страхование спецтехники')
        make('Haval H3', {'Альфа': (27000, 30000)})          # с франшизой — в режим «без франшизы» не идёт
        make('Haval F7', {'Пари': 999000})                    # тариф 99,9% — ошибка данных, отсекается

    def test_group_stats_and_cheapest(self):
        payload = tariffs.build_payload({})
        haval = next(row for row in payload['group_rows'] if row['group'] == 'Haval')
        self.assertEqual((haval['requests'], haval['offers']), (3, 9))
        self.assertAlmostEqual(haval['median'], 3.0)
        self.assertEqual(haval['cheapest']['company'], 'Зетта')
        self.assertFalse(haval['low_data'])

        gaz = next(row for row in payload['group_rows'] if row['group'] == 'ГАЗ')
        self.assertTrue(gaz['low_data'])
        self.assertIsNone(gaz['cheapest'])  # одна заявка — «дешевле всех» не назначаем

    def test_special_equipment_grouped_by_machine_kind(self):
        groups = {(row['dimension'], row['group']) for row in tariffs.build_payload({})['group_rows']}
        self.assertIn(('kind', 'Экскаватор'), groups)
        self.assertNotIn(('brand', 'SANY'), groups)

    def test_franchise_mode_and_outliers(self):
        franchise = tariffs.build_payload({'mode': tariffs.MODE_FRANCHISE})
        self.assertEqual(franchise['kpi']['offers'], 1)
        self.assertAlmostEqual(franchise['kpi']['median'], 2.7)
        self.assertFalse(any(point.tariff > tariffs.TARIFF_MAX for point in tariffs.collect_points({})))

    def test_heatmap_colors_relative_to_group(self):
        heatmap = tariffs.build_payload({})['heatmap']
        row = next(row for row in heatmap['rows'] if row['group'] == 'Haval')
        cells = {cell['company']: cell for cell in row['cells']}
        self.assertLess(cells['Зетта']['ratio'], 0.85)
        self.assertGreater(cells['Ингосстрах']['ratio'], 1.15)

    def test_filters(self):
        self.assertEqual(tariffs.build_payload({'object_class': 'lcv'})['kpi']['requests'], 1)
        self.assertEqual(tariffs.build_payload({'condition': 'used'})['kpi']['offers'], 0)

    def test_group_page_payload(self):
        payload = tariffs.build_group_payload({}, 'brand', 'Haval')
        self.assertEqual([row['company'] for row in payload['company_rows']], ['Зетта', 'Альфа', 'Ингосстрах'])
        self.assertEqual(payload['company_rows'][0]['wins'], 1)
        self.assertEqual(payload['model_rows'][0]['label'], 'M6')

    def test_hint_uses_brand_then_falls_back_to_class(self):
        haval = InsuranceSummary.objects.filter(request__vehicle_info='Автомобиль Haval M6').first()
        hint = tariffs.hint_for_request(haval.request, exclude_summary_id=haval.pk)
        self.assertEqual(hint['level'], 'class')  # без текущего свода по Haval остаётся 2 заявки из 3
        extra = make('Haval Jolion', {'Зетта': 21000})
        hint = tariffs.hint_for_request(extra.request, exclude_summary_id=extra.pk)
        self.assertEqual((hint['level'], hint['label']), ('group', 'Haval'))
        self.assertEqual(hint['cheapest'][0]['company'], 'Зетта')


class TariffViewTests(TestCase):
    def setUp(self):
        admin = User.objects.create_user(username='tariff_admin', password='pwd')
        admin.groups.add(Group.objects.get_or_create(name='Администраторы')[0])
        self.client.login(username='tariff_admin', password='pwd')
        for _ in range(3):
            make('Haval M6', {'Зетта': 20000, 'Ингосстрах': 40000})

    def test_pages_render(self):
        response = self.client.get(reverse('summaries:analytics_tariffs'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'марка × страховая')
        group = self.client.get(reverse('summaries:analytics_tariff_group', args=['brand', 'Haval']))
        self.assertEqual(group.status_code, 200)
        self.assertContains(group, 'Страховые по цене')
        self.assertEqual(self.client.get(reverse('summaries:analytics_tariff_group', args=['x', 'Haval'])).status_code, 404)
        overview = self.client.get(reverse('summaries:analytics'))
        self.assertContains(overview, reverse('summaries:analytics_tariffs'))

    def test_export(self):
        response = self.client.get(reverse('summaries:export_analytics_tariffs'), {'mode': 'no_franchise'})
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Классы', 'Марки и виды машин', 'Группа × страховая', 'Предложения'])
        self.assertEqual(workbook['Предложения'].max_row, 5 + 6)

    def test_regular_user_has_no_access(self):
        user = User.objects.create_user(username='tariff_user', password='pwd')
        user.groups.add(Group.objects.get_or_create(name='Пользователи')[0])
        self.client.login(username='tariff_user', password='pwd')
        self.assertNotEqual(self.client.get(reverse('summaries:analytics_tariffs')).status_code, 200)
