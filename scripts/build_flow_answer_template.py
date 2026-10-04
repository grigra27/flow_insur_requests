"""Сборка шаблона ответа страховщика templates/flow_answer_template.xlsx.

Шаблон отдаётся страховщикам кнопкой «Скачать актуальный шаблон ответа» и
читается обратно ExcelResponseProcessor (summaries/services/excel_services.py).
Адреса, которые читает парсер, менять нельзя: в обороте у страховщиков
остаются старые копии файла.

    B2      — страховая компания (выпадающий список)
    F2      — примечание (объединённая F2:J2)
    B3      — территория страхования (объединённая B3:J3)
    A6:J10  — годы 1–5: A год, B СС, D/E/F вариант 1, H/I/J вариант 2

Запуск из корня проекта:
    python scripts/build_flow_answer_template.py [путь_к_xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "templates" / "flow_answer_template.xlsx"
LOGO = ROOT / "core" / "brand" / "logo_online_broker.png"

# Названия как в справочнике InsuranceCompany на проде; «другое» — последним.
COMPANIES = [
    "Абсолют", "Альфа", "ВСК", "Зетта", "Ингосстрах", "Пари", "ПСБ-страхование",
    "Ренессанс", "РЕСО", "Росгосстрах", "Совкомбанк СК", "Согаз", "Согласие",
    "Энергогарант", "Югория", "другое",
]
INSTALLMENT_VALUES = "1,2,3,4,6,12"  # = ExcelResponseProcessor.VALID_INSTALLMENT_VALUES
YEARS = range(1, 6)
FIRST_YEAR_ROW = 6

BRAND_RED = "BA122B"
BRAND_SLATE = "495E5F"
BRAND_GRAY = "838280"
INPUT_FILL = PatternFill("solid", fgColor="FFF2B3")
LOCKED_FILL = PatternFill("solid", fgColor="EEF1F1")
HEADER_FILL = PatternFill("solid", fgColor=BRAND_SLATE)
SUBHEADER_FILL = PatternFill("solid", fgColor="DCE3E3")
SAMPLE_FILL = PatternFill("solid", fgColor="F4F4F4")
SAMPLE_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")

FONT = "Arial"
THIN = Side(style="thin", color="B7C0C0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
UNLOCKED = Protection(locked=False)

MONEY = "#,##0"
MONEY_KOP = "#,##0.00"

# Колонки строки года: (буква, формат, подпись в шапке)
YEAR_COLUMNS = [
    ("B", MONEY, "Страховая сумма, ₽"),
    ("D", MONEY_KOP, "Премия, ₽"),
    ("E", MONEY, "Франшиза, ₽"),
    ("F", "0", "Платежей в год"),
    ("H", MONEY_KOP, "Премия, ₽"),
    ("I", MONEY, "Франшиза, ₽"),
    ("J", "0", "Платежей в год"),
]

INSTRUCTIONS = [
    "Выберите свою компанию в поле «Страховая компания» — из выпадающего списка.",
    "Для каждого предлагаемого года заполните страховую сумму и премию. "
    "Страховую сумму указывайте в каждой строке, даже если она не меняется по годам.",
    "Годы, которые вы не предлагаете, оставьте пустыми.",
    "Франшизу укажите суммой в рублях; без франшизы — 0.",
    "«Платежей в год»: 1 — единовременно, 2 — раз в полгода, 4 — поквартально, 12 — ежемесячно.",
    "Вариант 2 заполняйте, только если в запросе есть второй вариант (например, другая франшиза).",
    "Территорию и территориальные ограничения — в поле «Территория страхования»; прочие условия "
    "(осмотр, согласование СБ, риски РНПК и т. п.) — в «Примечании».",
    "Вводите только числа — без «руб.», пробелов-букв и формул. Строки и столбцы не добавляйте.",
]

SAMPLE_ROWS = [
    (1, 2_000_000, 100_000, 0, 1, 75_000, 50_000, 1),
    (2, 1_700_000, 85_000, 0, 1, 64_000, 50_000, 1),
    (3, 1_450_000, 72_500, 0, 1, 54_500, 50_000, 1),
]
SAMPLE_TERRITORY = "Российская Федерация, за исключением ДНР, ЛНР, Херсонской и Запорожской областей"


def font(size=10, bold=False, italic=False, color="000000"):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def put(ws, coord, value=None, *, fnt=None, fill=None, align=None, border=None, fmt=None, unlocked=False):
    cell = ws[coord]
    if value is not None:
        cell.value = value
    cell.font = fnt or font()
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    if fmt:
        cell.number_format = fmt
    if unlocked:
        cell.protection = UNLOCKED
    return cell


def box_range(ws, cells, **kwargs):
    for coord in cells:
        put(ws, coord, **kwargs)


def label(ws, coord, text):
    put(ws, coord, text, fnt=font(10, bold=True, color=BRAND_SLATE),
        align=Alignment(horizontal="left", vertical="center", wrap_text=True))


def table_header(ws, top, fill, sub_fill, text_color, variant_2_title):
    """Двухстрочная шапка таблицы лет: группы «Вариант 1/2» над подписями колонок."""
    hdr_font = font(10, bold=True, color=text_color)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, title in (("A", "Год страхования"), ("B", "Страховая сумма, ₽")):
        ws.merge_cells(f"{col}{top}:{col}{top + 1}")
        box_range(ws, [f"{col}{top}", f"{col}{top + 1}"], fill=fill, border=BOX, align=center, fnt=hdr_font)
        ws[f"{col}{top}"].value = title
    for first, last, title in (("D", "F", "Вариант 1"), ("H", "J", variant_2_title)):
        ws.merge_cells(f"{first}{top}:{last}{top}")
        box_range(ws, [f"{c}{top}" for c in "DEFHIJ" if first <= c <= last],
                  fill=fill, border=BOX, align=center, fnt=hdr_font)
        ws[f"{first}{top}"].value = title
    for col, _fmt, title in YEAR_COLUMNS[1:]:
        put(ws, f"{col}{top + 1}", title, fnt=font(9, bold=True, color=BRAND_SLATE),
            fill=sub_fill, border=BOX, align=center)


def add_validations(ws):
    last = FIRST_YEAR_ROW + len(YEARS) - 1

    company = DataValidation(
        type="list", formula1='"' + ",".join(COMPANIES) + '"', allow_blank=True,
        showErrorMessage=True, errorTitle="Страховая компания",
        error="Выберите компанию из выпадающего списка.",
        promptTitle="Страховая компания", prompt="Выберите свою компанию из списка.",
    )
    company.add("B2")

    def money(cols, operator, formula, prompt):
        dv = DataValidation(
            type="decimal", operator=operator, formula1=formula, allow_blank=True,
            showErrorMessage=True, errorTitle="Нужно число",
            error="Введите число без «руб.», букв и формул.",
            promptTitle="Сумма в рублях", prompt=prompt,
        )
        for col in cols:
            dv.add(f"{col}{FIRST_YEAR_ROW}:{col}{last}")
        return dv

    sums = money("B", "greaterThan", "0",
                 "Страховая сумма этого года. Заполните для каждого года, даже если она не меняется.")
    premiums = money("DH", "greaterThan", "0", "Страховая премия за этот год.")
    franchises = money("EI", "greaterThanOrEqual", "0", "Франшиза в рублях; без франшизы — 0.")

    installment = DataValidation(
        type="list", formula1=f'"{INSTALLMENT_VALUES}"', allow_blank=True,
        showErrorMessage=True, errorTitle="Платежей в год",
        error="Допустимо: 1, 2, 3, 4, 6 или 12.",
        promptTitle="Платежей в год",
        prompt="1 — единовременно, 2 — раз в полгода, 4 — поквартально, 12 — ежемесячно.",
    )
    for col in "FJ":
        installment.add(f"{col}{FIRST_YEAR_ROW}:{col}{last}")

    for dv in (company, sums, premiums, franchises, installment):
        ws.add_data_validation(dv)


def build(output: Path = OUTPUT) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ответ СК"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100

    widths = {"A": 13, "B": 22, "C": 1.2, "D": 14, "E": 13, "F": 11,
              "G": 1.2, "H": 14, "I": 13, "J": 11}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    # 1. Заголовок и логотип
    ws.row_dimensions[1].height = 42
    ws.merge_cells("A1:G1")
    put(ws, "A1", "Ответ страховой компании на запрос котировки",
        fnt=font(15, bold=True, color=BRAND_RED), align=Alignment(vertical="center"))
    if LOGO.is_file():
        logo = XLImage(str(LOGO))
        logo.height = 46
        logo.width = int(46 * 600 / 250)
        ws.add_image(logo, "I1")

    # 2. Компания и примечание
    ws.row_dimensions[2].height = 58
    label(ws, "A2", "Страховая компания")
    put(ws, "B2", fnt=font(11, bold=True), fill=INPUT_FILL, border=BOX,
        align=Alignment(vertical="center"), unlocked=True)
    ws.merge_cells("D2:E2")
    put(ws, "D2", "Примечание", fnt=font(10, bold=True, color=BRAND_SLATE),
        align=Alignment(horizontal="right", vertical="top", indent=1))
    ws.merge_cells("F2:J2")
    box_range(ws, [f"{c}2" for c in "FGHIJ"], fill=INPUT_FILL, border=BOX,
              align=Alignment(vertical="top", wrap_text=True), unlocked=True)

    # 3. Территория
    ws.row_dimensions[3].height = 46
    label(ws, "A3", "Территория страхования")
    ws.merge_cells("B3:J3")
    box_range(ws, [f"{c}3" for c in "BCDEFGHIJ"], fill=INPUT_FILL, border=BOX,
              align=Alignment(vertical="top", wrap_text=True), unlocked=True)

    # 4–10. Таблица лет
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 26
    table_header(ws, 4, HEADER_FILL, SUBHEADER_FILL, "FFFFFF", "Вариант 2 (если запрошен)")
    for year in YEARS:
        row = FIRST_YEAR_ROW + year - 1
        ws.row_dimensions[row].height = 20
        put(ws, f"A{row}", year, fnt=font(10, bold=True, color=BRAND_SLATE), fill=LOCKED_FILL,
            border=BOX, align=Alignment(horizontal="center", vertical="center"))
        for col, fmt, _title in YEAR_COLUMNS:
            put(ws, f"{col}{row}", fill=INPUT_FILL, border=BOX, fmt=fmt,
                align=Alignment(horizontal="right" if fmt != "0" else "center", vertical="center"),
                unlocked=True)

    # 12–20. Как заполнить
    top = 12
    ws.merge_cells(f"A{top}:J{top}")
    put(ws, f"A{top}", "Как заполнить", fnt=font(11, bold=True, color=BRAND_RED),
        align=Alignment(vertical="bottom"))
    ws.row_dimensions[top].height = 22
    for idx, text in enumerate(INSTRUCTIONS, start=1):
        row = top + idx
        ws.merge_cells(f"B{row}:J{row}")
        put(ws, f"A{row}", idx, fnt=font(10, bold=True, color=BRAND_RED),
            align=Alignment(horizontal="right", vertical="top", indent=1))
        put(ws, f"B{row}", text, fnt=font(10), align=Alignment(vertical="top", wrap_text=True))
        ws.row_dimensions[row].height = 28 if len(text) > 95 else 15

    # Образец заполнения
    top = top + len(INSTRUCTIONS) + 2
    ws.merge_cells(f"A{top}:J{top}")
    put(ws, f"A{top}", "Образец заполнения (не заполняйте — только пример)",
        fnt=font(11, bold=True, color=BRAND_GRAY), align=Alignment(vertical="bottom"))
    ws.row_dimensions[top].height = 22
    sample_top = top + 1
    ws.row_dimensions[sample_top].height = 20
    ws.row_dimensions[sample_top + 1].height = 26
    table_header(ws, sample_top, SAMPLE_HEADER_FILL, SAMPLE_FILL, "404040", "Вариант 2")
    gray = font(10, color="595959")
    for offset, values in enumerate(SAMPLE_ROWS):
        row = sample_top + 2 + offset
        ws.row_dimensions[row].height = 18
        for col, value in zip("ABDEFHIJ", values):
            fmt = next((f for c, f, _t in YEAR_COLUMNS if c == col), "0")
            put(ws, f"{col}{row}", value, fnt=gray, fill=SAMPLE_FILL, border=BOX, fmt=fmt,
                align=Alignment(horizontal="center" if col in "AFJ" else "right", vertical="center"))
    row = sample_top + 2 + len(SAMPLE_ROWS)
    ws.row_dimensions[row].height = 18
    put(ws, f"A{row}", "Территория", fnt=font(9, bold=True, color="595959"),
        align=Alignment(vertical="center"))
    ws.merge_cells(f"B{row}:J{row}")
    put(ws, f"B{row}", SAMPLE_TERRITORY, fnt=gray, fill=SAMPLE_FILL, border=BOX,
        align=Alignment(vertical="center"))

    add_validations(ws)

    # Защита без пароля: редактируются только жёлтые ячейки, высоту строк менять можно.
    ws.protection.sheet = True
    ws.protection.formatRows = False
    ws.protection.formatColumns = False
    ws.protection.formatCells = False

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:J{row}"
    ws.sheet_view.selection[0].activeCell = "B2"
    ws.sheet_view.selection[0].sqref = "B2"

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT
    print(build(target))
