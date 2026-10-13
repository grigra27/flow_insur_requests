"""Готовые значения формул в xlsx, собранном openpyxl.

openpyxl записывает в ячейку только формулу (``<f>…</f><v></v>``), без результата. Excel пересчитывает
формулы при открытии, но файл, скачанный из интернета, открывается в «защищённом просмотре» — там
пересчёта нет, и ИТОГО/тарифы выглядят пустыми, пока не нажать «Разрешить редактирование».

with_cached_values() считает формулы по числам, уже записанным в файл, и дописывает результат рядом
с формулой, как это делает сам Excel при сохранении. Формулы остаются живыми (при правке премии ИТОГО
пересчитается), а в защищённом просмотре сразу видны значения.

Поддерживается то, что используют своды: числа, строки, ссылки и диапазоны на том же листе, + - * /,
сравнения, SUM, IF, AND, OR. Формулу, которую посчитать не удалось, оставляем без значения — Excel
посчитает её сам после «Разрешить редактирование».
"""
from __future__ import annotations

import logging
import re
import zipfile
from io import BytesIO
from typing import Dict, Optional, Tuple
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

logger = logging.getLogger(__name__)

TOKEN = re.compile(r'''\s*(?:
    (?P<num>\d+(?:\.\d+)?)|
    (?P<str>"(?:[^"]|"")*")|
    (?P<range>\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+)|
    (?P<ref>\$?[A-Z]{1,3}\$?\d+)|
    (?P<func>[A-Z]+)\s*\(|
    (?P<op><>|<=|>=|[-+*/=<>(),])
)''', re.VERBOSE)


class Unsupported(Exception):
    """Формулу посчитать нельзя — оставляем её Excel."""


def _tokens(formula: str):
    pos, out = 0, []
    text = formula.lstrip('=')
    while pos < len(text):
        match = TOKEN.match(text, pos)
        if not match or match.end() == pos:
            raise Unsupported(formula)
        kind = match.lastgroup
        out.append((kind, match.group(kind)))
        pos = match.end()
        while pos < len(text) and text[pos].isspace():
            pos += 1
    return out


def _number(value) -> float:
    if value is None or value == '':
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    raise Unsupported(f'не число: {value!r}')


class _Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.cache: Dict[str, object] = {}

    def value(self, coord: str):
        coord = coord.replace('$', '')
        if coord in self.cache:
            return self.cache[coord]
        raw = self.ws[coord].value
        if isinstance(raw, str) and raw.startswith('='):
            self.cache[coord] = None  # защита от циклов
            result = _Parser(_tokens(raw), self).parse()
            self.cache[coord] = result
            return result
        return raw

    def range_values(self, ref: str):
        start, end = ref.replace('$', '').split(':')
        c1, r1 = re.match(r'([A-Z]+)(\d+)', start).groups()
        c2, r2 = re.match(r'([A-Z]+)(\d+)', end).groups()
        for col in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
            for row in range(int(r1), int(r2) + 1):
                yield self.value(f'{get_column_letter(col)}{row}')


class _Parser:
    """Разбор с приоритетами: сравнение < сложение < умножение < унарный минус < операнд."""

    def __init__(self, tokens, sheet: _Sheet):
        self.tokens, self.pos, self.sheet = tokens, 0, sheet

    def peek(self) -> Tuple[Optional[str], Optional[str]]:
        return self.tokens[self.pos] if self.pos < len(self.tokens) else (None, None)

    def take(self, value=None):
        kind, text = self.peek()
        if kind is None or (value is not None and text != value):
            raise Unsupported(f'ожидалось {value!r}')
        self.pos += 1
        return kind, text

    def parse(self):
        result = self.comparison()
        if self.pos != len(self.tokens):
            raise Unsupported('лишние символы')
        return result

    def comparison(self):
        left = self.additive()
        kind, text = self.peek()
        if kind == 'op' and text in ('=', '<>', '<', '>', '<=', '>='):
            self.take()
            right = self.additive()
            if text in ('=', '<>'):
                if isinstance(left, str) or isinstance(right, str):
                    equal = str(left if left is not None else '') == str(right if right is not None else '')
                else:
                    equal = _number(left) == _number(right)
                return equal if text == '=' else not equal
            a, b = _number(left), _number(right)
            return {'<': a < b, '>': a > b, '<=': a <= b, '>=': a >= b}[text]
        return left

    def additive(self):
        value = self.term()
        while self.peek() in (('op', '+'), ('op', '-')):
            _, op = self.take()
            right = self.term()
            value = _number(value) + _number(right) if op == '+' else _number(value) - _number(right)
        return value

    def term(self):
        value = self.unary()
        while self.peek() in (('op', '*'), ('op', '/')):
            _, op = self.take()
            right = self.unary()
            if op == '*':
                value = _number(value) * _number(right)
            else:
                divisor = _number(right)
                if divisor == 0:
                    raise Unsupported('деление на ноль')
                value = _number(value) / divisor
        return value

    def unary(self):
        if self.peek() == ('op', '-'):
            self.take()
            return -_number(self.unary())
        return self.operand()

    def operand(self):
        kind, text = self.take()
        if kind == 'num':
            return float(text)
        if kind == 'str':
            return text[1:-1].replace('""', '"')
        if kind == 'ref':
            return self.sheet.value(text)
        if kind == 'op' and text == '(':
            value = self.comparison()
            self.take(')')
            return value
        if kind == 'func':
            return self.function(text)
        raise Unsupported(f'неожиданный элемент {text!r}')

    def arguments(self):
        args = []
        if self.peek() == ('op', ')'):
            self.take(')')
            return args
        while True:
            kind, text = self.peek()
            if kind == 'range':
                self.take()
                args.append(list(self.sheet.range_values(text)))
            else:
                args.append(self.comparison())
            kind, text = self.take()
            if text == ')':
                return args
            if text != ',':
                raise Unsupported('ожидалась запятая')

    def function(self, name: str):
        if name == 'IF':
            args = self.arguments()
            if len(args) not in (2, 3):
                raise Unsupported('IF')
            condition = args[0]
            truthy = bool(_number(condition)) if not isinstance(condition, str) else bool(condition)
            return args[1] if truthy else (args[2] if len(args) == 3 else False)
        args = self.arguments()
        flat = [v for arg in args for v in (arg if isinstance(arg, list) else [arg])]
        if name == 'SUM':
            return sum(_number(v) for v in flat if not isinstance(v, str))
        if name == 'AND':
            return all(bool(_number(v)) for v in flat)
        if name == 'OR':
            return any(bool(_number(v)) for v in flat)
        raise Unsupported(f'функция {name}')


def _xml_value(value) -> Optional[Tuple[str, bool]]:
    """(текст для <v>, строковый ли результат) или None, если значение записывать не нужно."""
    if value is None:
        return None
    if isinstance(value, bool):
        return ('1' if value else '0'), False
    if isinstance(value, (int, float)):
        number = float(value)
        return (str(int(number)) if number.is_integer() else repr(number)), False
    return escape(str(value)), True


def with_cached_values(data: bytes) -> bytes:
    """xlsx (собранный openpyxl) → тот же xlsx, где у формул записаны готовые значения."""
    try:
        wb = load_workbook(BytesIO(data))
    except Exception:  # noqa: BLE001 — файл не наш, отдаём как есть
        logger.exception('Cached values: cannot read workbook')
        return data

    values: Dict[int, Dict[str, Tuple[str, bool]]] = {}
    for index, ws in enumerate(wb.worksheets, start=1):
        sheet = _Sheet(ws)
        found = {}
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    try:
                        xml = _xml_value(sheet.value(cell.coordinate))
                    except Unsupported:
                        continue
                    if xml is not None:
                        found[cell.coordinate] = xml
        if found:
            values[index] = found
    if not values:
        return data

    source = zipfile.ZipFile(BytesIO(data))
    output = BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            match = re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml', item.filename)
            if match and int(match.group(1)) in values:
                content = _patch_sheet(content.decode('utf-8'), values[int(match.group(1))]).encode('utf-8')
            target.writestr(item, content)
    return output.getvalue()


CELL_WITH_FORMULA = re.compile(r'<c r="(?P<coord>[A-Z]+\d+)"(?P<attrs>[^>]*)>(?P<f><f>.*?</f>)<v\s*/>|'
                               r'<c r="(?P<coord2>[A-Z]+\d+)"(?P<attrs2>[^>]*)>(?P<f2><f>.*?</f>)<v></v>', re.S)


def _patch_sheet(xml: str, values: Dict[str, Tuple[str, bool]]) -> str:
    def replace(match):
        coord = match.group('coord') or match.group('coord2')
        attrs = match.group('attrs') if match.group('coord') else match.group('attrs2')
        formula = match.group('f') or match.group('f2')
        if coord not in values:
            return match.group(0)
        text, is_string = values[coord]
        attrs = re.sub(r'\s+t="[^"]*"', '', attrs)
        if is_string:
            attrs += ' t="str"'
        return f'<c r="{coord}"{attrs}>{formula}<v>{text}</v>'

    return CELL_WITH_FORMULA.sub(replace, xml)
