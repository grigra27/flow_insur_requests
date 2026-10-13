"""Готовые значения формул в xlsx: видны без пересчёта (как в «защищённом просмотре» Excel)."""
from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from core.excel_cached_values import with_cached_values


def _cached(wb):
    buffer = BytesIO()
    wb.save(buffer)
    data = with_cached_values(buffer.getvalue())
    return load_workbook(BytesIO(data), data_only=True), load_workbook(BytesIO(data))


class CachedValuesTests(SimpleTestCase):
    def test_sum_and_rate_formulas_get_values_and_stay_formulas(self):
        wb = Workbook()
        ws = wb.active
        ws['C10'], ws['E10'], ws['E11'] = 7_300_000, 98_690, 98_690.5
        ws['C11'] = 0
        ws['D10'] = '=IF(AND(C10<>0,E10<>0),E10/C10,"")'
        ws['D11'] = '=IF(AND(C11<>0,E11<>0),E11/C11,"")'  # сумма 0 — пустая строка, как в Excel
        ws['H10'] = '=SUM(E10:E11)'
        ws['H12'] = '=SUM(E10,E11)-1'
        values, formulas = _cached(wb)
        sheet = values.active
        self.assertAlmostEqual(sheet['D10'].value, 98_690 / 7_300_000)
        self.assertIn(sheet['D11'].value, ('', None))  # пустая строка — Excel покажет пустую ячейку
        self.assertEqual(sheet['H10'].value, 197_380.5)
        self.assertEqual(sheet['H12'].value, 197_379.5)
        self.assertEqual(formulas.active['H10'].value, '=SUM(E10:E11)')  # формула на месте

    def test_formula_on_formula_and_blank_cells(self):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = '=SUM(B1:B3)'  # B2 пустая
        ws['B1'], ws['B3'] = 10, 5
        ws['A2'] = '=A1*2'
        values, _ = _cached(wb)
        self.assertEqual((values.active['A1'].value, values.active['A2'].value), (15, 30))

    def test_unsupported_formula_is_left_to_excel(self):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = '=VLOOKUP(1,B1:C2,2,FALSE)'
        ws['A2'] = '=1/0'
        ws['A3'] = '=2+2'
        values, formulas = _cached(wb)
        self.assertIsNone(values.active['A1'].value)
        self.assertIsNone(values.active['A2'].value)
        self.assertEqual(values.active['A3'].value, 4)
        self.assertEqual(formulas.active['A1'].value, '=VLOOKUP(1,B1:C2,2,FALSE)')

    def test_every_sheet_is_patched(self):
        wb = Workbook()
        wb.active['A1'] = '=1+1'
        wb.create_sheet('Второй')['B2'] = '=3*3'
        values, _ = _cached(wb)
        self.assertEqual(values['Sheet']['A1'].value, 2)
        self.assertEqual(values['Второй']['B2'].value, 9)
